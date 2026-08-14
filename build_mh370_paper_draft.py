#!/usr/bin/env python3
"""Build the MH370 sections 1--2 working manuscript and companion evidence memo."""

from __future__ import annotations

import math
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "paper_draft"
FIG = OUT / "figures"
OUT.mkdir(parents=True, exist_ok=True)
FIG.mkdir(parents=True, exist_ok=True)


SATCOM = [
    ("ATSB wide area (2014)", 27.70, 16.33, 39.07),
    ("ATSB priority area (2014)", 29.80, 27.45, 32.15),
    ("ATSB medium area (2014)", 29.60, 24.50, 34.70),
    ("ATSB update (Oct. 2014)", 35.90, 32.18, 39.62),
    ("Ashton et al. (2014)", 34.70, 33.72, 35.68),
    ("Anderson et al. (2014)", 37.60, 36.62, 38.58),
    ("Spinor (2014)", 38.50, 37.52, 39.48),
    ("Ulich (2014)", 40.20, 40.00, 40.40),
    ("Martin (2014)", 40.55, 39.57, 41.53),
    ("Pleter et al. (2015)", 37.50, 36.13, 38.87),
    ("Fah (2015)", 37.50, 34.56, 40.44),
    ("Davey et al. (2015)", 38.00, 35.06, 40.94),
    ("GlobusMax (2015)", 40.00, 39.02, 40.98),
    ("ATSB (2015)", 37.70, 35.74, 39.66),
    ("ATSB (2016)", 34.30, 32.54, 36.06),
    ("Iannello & Godfrey (2016)", 26.90, 25.92, 27.88),
    ("SK999 (2016)", 29.50, 22.05, 36.95),
    ("Iannello (2017)", 34.00, 28.51, 39.49),
    ("Godfrey (2017)", 30.00, 29.02, 30.98),
    ("Nederland (2017)", 31.10, 30.12, 32.08),
    ("Gilbert (2017)", 33.50, 31.93, 35.07),
    ("Marchand et al. (2018)", 12.00, 11.02, 12.98),
    ("Ulich (2018)", 31.60, 30.62, 32.58),
    ("Iannello (2018)", 22.00, 21.02, 22.98),
    ("Kristensen (2018)", 13.30, 11.73, 14.87),
    ("Kristensen, 2nd minimum", 34.60, 33.03, 36.17),
]

DRIFT = [
    ("Pattiaratchi & Wijeratne", 30.50, 28.05, 32.95),
    ("Rydberg", 25.00, 12.26, 37.74),
    ("Daniel", 30.00, 25.10, 34.90),
    ("Jansen et al.", 31.50, 28.07, 34.93),
    ("Durgadoo et al.", 23.50, 14.19, 32.81),
    ("Trinanes et al.", 32.00, 26.12, 37.88),
    ("Griffin, Oke & Jones", 32.00, 26.12, 37.88),
    ("Godfrey (2017)", 28.00, 19.18, 36.82),
    ("Nesterov", 30.00, 25.10, 34.90),
    ("Godfrey (2018)", 27.00, 19.16, 34.84),
    ("Griffin & Oke", 30.50, 25.11, 35.89),
    ("Miron et al.", 25.00, 17.16, 32.84),
]

PHR4_RATING5 = [
    (2, -34.45100, 91.355206, 70),
    (3, -34.45289, 91.354443, 23),
    (4, -34.45283, 91.355169, 43),
    (5, -34.45326, 91.355020, 28),
    (6, -34.45946, 91.364210, 25),
    (18, -34.49835, 91.306980, 69),
    (19, -34.50437, 91.305206, 41),
    (26, -34.54438, 91.304036, 49),
    (27, -34.54901, 91.292308, 39),
]


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text))


def make_evidence_flow(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.0, 3.75))
    ax.set_xlim(0, 7.0)
    ax.set_ylim(0, 3.75)
    ax.axis("off")

    colours = {
        "obs": "#E8F1F8",
        "state": "#D8EAD2",
        "transition": "#F8E7C2",
        "conditional": "#EFE3F5",
        "search": "#E8E8E8",
        "edge": "#23395B",
    }

    def box(x, y, w, h, text, colour, fs=8.2, ls="-"):
        patch = FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.025,rounding_size=0.06",
            facecolor=colour, edgecolor=colours["edge"], linewidth=1.15,
            linestyle=ls,
        )
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs)
        return patch

    def arrow(x1, y1, x2, y2, style="-", rad=0.0):
        a = FancyArrowPatch(
            (x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=11,
            linewidth=1.1, color=colours["edge"], linestyle=style,
            connectionstyle=f"arc3,rad={rad}",
        )
        ax.add_patch(a)

    ax.text(0.12, 3.48, "Observation-level integration and scenario separation", fontsize=11.8, weight="semibold")
    box(0.12, 2.02, 1.48, 1.02,
        "Inputs through 00:11\nradar state; BTO/BFO\nand RSL; performance,\nfuel, wind, control priors",
        colours["obs"], fs=7.0)
    box(1.91, 2.13, 1.25, 0.80,
        r"00:11 airborne" "\n" r"posterior $p(x_6|y)$",
        colours["state"], fs=8.0)
    box(3.52, 2.02, 1.50, 1.02,
        "Terminal-flight kernel\ncontrolled descent,\nglide and uncontrolled\nfamilies",
        colours["transition"], fs=7.1)
    box(5.43, 2.13, 1.30, 0.80,
        r"Impact posterior" "\n" r"$p(x_I|y)$",
        colours["state"], fs=8.2)

    box(1.91, 0.70, 1.25, 0.68, "00:19 observations\nEOF sensitivity only", "white", fs=7.3, ls="--")
    box(3.52, 0.70, 1.50, 0.68, "Ocean-drift likelihood", colours["conditional"], fs=7.6)
    box(5.43, 0.88, 1.30, 0.58, "Pleiades likelihood\nconditional branch", colours["conditional"], fs=7.0)
    box(5.43, 0.18, 1.30, 0.48, "Search non-detection", colours["search"], fs=6.6)

    arrow(1.60, 2.53, 1.91, 2.53)
    arrow(3.16, 2.53, 3.52, 2.53)
    arrow(5.02, 2.53, 5.43, 2.53)
    arrow(2.54, 2.13, 2.54, 1.38, style="--")
    arrow(4.27, 1.38, 5.67, 2.13)
    arrow(6.08, 1.46, 6.08, 2.13)
    arrow(6.08, 0.66, 6.08, 2.13, rad=-0.12)

    fig.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_diamond(ax, centre, low, high, y, height=0.32, colour="#B23A48"):
    ax.add_patch(Polygon([[low, y], [centre, y + height], [high, y], [centre, y - height]], closed=True,
                         facecolor=colour, edgecolor="black", linewidth=0.7, zorder=4))


def draw_forest(ax, data, title, pooled, qtext, arithmetic, xlim=(9, 43), fontsize=7.3):
    n = len(data)
    ypos = np.arange(n, 0, -1)
    centres = np.array([x[1] for x in data])
    lows = np.array([x[2] for x in data])
    highs = np.array([x[3] for x in data])
    ax.hlines(ypos, lows, highs, color="#4D4D4D", linewidth=0.85, zorder=1)
    ax.scatter(centres, ypos, s=16, marker="s", color="#1F4E79", edgecolor="white", linewidth=0.3, zorder=3)
    ax.set_yticks(ypos)
    ax.set_yticklabels([x[0] for x in data], fontsize=fontsize)
    ax.tick_params(axis="y", length=0, pad=4)
    ax.set_xlim(*xlim)
    ax.set_ylim(-1.2, n + 1.4)
    ax.grid(axis="x", color="#D8D8D8", linewidth=0.6)
    ax.axvline(arithmetic, color="#777777", linestyle="--", linewidth=0.8, label="Arithmetic mean")
    pc, pl, ph = pooled
    add_diamond(ax, pc, pl, ph, 0, height=0.32)
    ax.text(xlim[0] + 0.2, 0, "Pooled", va="center", fontsize=fontsize, weight="semibold")
    ax.text(ph + 0.5, 0, f"{pc:.2f} [{pl:.2f}, {ph:.2f}]", va="center", fontsize=fontsize)
    ax.set_title(title, loc="left", fontsize=11.5, pad=6, weight="semibold")
    ax.text(xlim[1], n + 0.8, qtext, ha="right", va="center", fontsize=fontsize, color="#333333")
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.set_xlabel("Estimated latitude (degrees south)", fontsize=8.5)
    ax.tick_params(axis="x", labelsize=7.5)


def make_forest(path: Path) -> None:
    fig = plt.figure(figsize=(7.2, 9.45))
    gs = fig.add_gridspec(2, 1, height_ratios=[2.05, 1], hspace=0.32)
    ax1 = fig.add_subplot(gs[0])
    ax2 = fig.add_subplot(gs[1])
    draw_forest(
        ax1, SATCOM, "(a) Satellite-communications estimates (n = 26)",
        (33.34, 29.13, 37.54), r"Random-effects: $Q=6368.5$, df = 25, $p<0.0001$",
        arithmetic=32.2327, fontsize=7.5,
    )
    draw_forest(
        ax2, DRIFT, "(b) Ocean-drift estimates (n = 12)",
        (30.17, 28.74, 31.60), r"Fixed-effect descriptive pool: $Q=6.55$, df = 11, $p=0.83$",
        arithmetic=28.75, fontsize=8.0,
    )
    fig.suptitle("Published arc-latitude estimates synthesised in Large (2019)", fontsize=13.5, y=0.995)
    fig.text(0.5, 0.005, "Intervals include values reconstructed or imputed from published ranges; they are descriptive and are not independent measurement likelihoods.", ha="center", fontsize=7.8)
    fig.savefig(path, dpi=330, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_holland_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(7.1, 3.85))
    ax.set_xlim(0, 7.1); ax.set_ylim(0, 3.85); ax.axis("off")
    cols = ["#E8F1F8", "#F8E7C2", "#F3D5D7"]
    labels = [
        (0.08, 2.78, 2.00, 0.84, "Recorded facts", "182 Hz at 00:19:29\n−2 Hz at 00:19:37\ntransient R600/R1200"),
        (2.54, 2.78, 2.00, 0.84, "Signal-model inference", "After track, bias and OCXO\nassumptions: downward\nline-of-sight velocity"),
        (5.00, 2.78, 2.00, 0.84, "Dynamic interpretation", "Consistent with rapid descent;\ndoes not identify control\nor impact distance"),
    ]
    for i,(x,y,w,h,head,body) in enumerate(labels):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.025",facecolor=cols[i],edgecolor="#23395B",linewidth=1.15))
        ax.text(x+w/2,y+h-0.15,head,ha="center",va="top",fontsize=8.5,weight="semibold")
        ax.text(x+w/2,y+0.30,body,ha="center",va="center",fontsize=7.5,linespacing=1.15)
    for x in [2.08, 4.54]:
        ax.add_patch(FancyArrowPatch((x,3.20),(x+0.46,3.20),arrowstyle="-|>",mutation_scale=11,color="#23395B",linewidth=1.1))

    caveats = [
        "OCXO correction is extrapolated from seven, mostly much longer, outages.",
        "The 4.3-Hz residual spread mixes fast noise with slow/intermittent bias.",
        "Independent bounds and their midpoints are not a coherent trajectory posterior.",
        "ATSB simulations were uncontrolled and did not cover every possible condition.",
    ]
    ax.text(0.10,2.43,"Where additional assumptions enter",fontsize=9.7,weight="semibold")
    for i,t in enumerate(caveats):
        y=2.06-i*0.36
        ax.add_patch(Rectangle((0.14,y-0.055),0.11,0.11,facecolor="#B23A48",edgecolor="none"))
        ax.text(0.36,y,t,va="center",fontsize=7.7)
    ax.add_patch(Rectangle((0.08,0.15),6.92,0.40,facecolor="#EEF2F5",edgecolor="#6B7C8C",linewidth=0.8))
    ax.text(0.18,0.35,
            "Conditional conclusion: strong evidence for descent at two instants; weak identification of what happened next.",
            va="center",fontsize=7.8)
    fig.savefig(path,dpi=300,bbox_inches="tight",facecolor="white")
    plt.close(fig)


def make_radar_figure(path: Path) -> None:
    times = np.array([17+56/60, 18+1/60+59/3600, 18+2/60, 18+15/60+25/3600, 18+22/60+12/3600])
    alts = np.array([44700, 58200, 4800, 29500, 29500])
    labels = ["near Penang\n44,700 ft", "58,200 ft", "Pulau Perak\n4,800 ft", "reappearance\n29,500 ft", "final return\n29,500 ft"]
    fig, ax = plt.subplots(figsize=(7.2, 3.65))
    ax.plot(times, alts/1000, color="#B23A48", linewidth=1.6, marker="o", markersize=6)
    for x,y,l in zip(times,alts/1000,labels):
        ax.annotate(l,(x,y),xytext=(0,9),textcoords="offset points",ha="center",fontsize=7.4)
    ax.axvspan(18+3/60+9/3600,18+15/60+25/3600,color="#D9D9D9",alpha=0.55,label="reported radar gap")
    ax.set_ylim(0,67); ax.set_xlim(times.min()-0.02,times.max()+0.02)
    tick_times=np.array([17+56/60,18,18+4/60,18+8/60,18+12/60,18+16/60,18+20/60])
    ax.set_xticks(tick_times); ax.set_xticklabels([f"{int(t):02d}:{int(round((t%1)*60)):02d}" for t in tick_times])
    ax.set_ylabel("Reported height label (thousand ft)",fontsize=8.5); ax.set_xlabel("UTC on 7 March 2014",fontsize=8.5)
    ax.tick_params(labelsize=7.5)
    ax.grid(color="#DDDDDD",linewidth=0.7)
    ax.set_title("Unusable height labels; a distinct contact gap",loc="left",fontsize=11.5)
    ax.annotate("reported loss\n18:03:09",xy=(18+3/60+9/3600,31),xytext=(17.98,38),
                arrowprops=dict(arrowstyle="->",color="#555555",lw=0.9),fontsize=7.6,ha="center")
    ax.annotate("reported reacquisition\n18:15:25",xy=(18+15/60+25/3600,29.5),xytext=(18.31,38),
                arrowprops=dict(arrowstyle="->",color="#555555",lw=0.9),fontsize=7.6,ha="center")
    ax.legend(loc="upper right",fontsize=7.3)
    fig.subplots_adjust(left=0.11,right=0.985,top=0.86,bottom=0.23)
    fig.text(0.5,0.04,"Official assessment: altitude/speed extraction was inherently unreliable; latitude/longitude was reasonably reliable.",ha="center",fontsize=7.2)
    fig.savefig(path,dpi=300,facecolor="white")
    plt.close(fig)


def make_pleiades_figure(path: Path) -> None:
    fig = plt.figure(figsize=(7.2, 3.85))
    gs = fig.add_gridspec(1, 2, width_ratios=[1.0, 1.12], wspace=0.30,left=0.12,right=0.98,top=0.88,bottom=0.18)
    ax = fig.add_subplot(gs[0])
    ids=np.array([r[0] for r in PHR4_RATING5]); lat=np.array([r[1] for r in PHR4_RATING5]); lon=np.array([r[2] for r in PHR4_RATING5]); area=np.array([r[3] for r in PHR4_RATING5])
    sizes=area*3.0
    ax.scatter(lon,lat,s=sizes,color="#5B8E7D",edgecolor="#16324F",alpha=0.78)
    for i,x,y in zip(ids,lon,lat):
        if int(i) in {2,3,4,5}: continue
        offsets={6:(5,2),18:(5,-11),19:(5,5),26:(5,5),27:(5,0)}
        ax.annotate(str(i),(x,y),xytext=offsets[int(i)],textcoords="offset points",fontsize=7.4)
    ax.annotate("2–5",(lon[:4].mean(),lat[:4].mean()),xytext=(-7,11),textcoords="offset points",fontsize=7.4,ha="center")
    ax.set_xlabel("Longitude (°E)",fontsize=8.2); ax.set_ylabel("Latitude (°)",fontsize=8.2)
    ax.tick_params(labelsize=7.3)
    ax.set_title("PHR_4 north-east scene\n9 ‘probably man-made’ objects",loc="left",fontsize=10.2)
    ax.grid(color="#DDDDDD",linewidth=0.6)
    ax.invert_yaxis()
    ax.text(0.0,-0.18,"Marker area ∝ reported bright footprint (23–70 m²).",transform=ax.transAxes,fontsize=7.1,clip_on=False)

    ax2=fig.add_subplot(gs[1]); ax2.axis("off"); ax2.set_xlim(0,10); ax2.set_ylim(0,10)
    ax2.set_title("What 0.5-m sampling permits",loc="left",fontsize=10.2)
    # Pixel grid with a schematic 2-m feature.
    for i in range(10):
        for j in range(5):
            fc="#E7EEF4"
            if j==2 and 3<=i<=6: fc="#1F4E79"
            ax2.add_patch(Rectangle((0.7+i*0.50,6.8+j*0.45),0.50,0.45,facecolor=fc,edgecolor="white",linewidth=0.6))
    ax2.annotate("~2 m = ~4 pixels",xy=(3.2,7.92),xytext=(3.2,9.55),ha="center",arrowprops=dict(arrowstyle="-[,widthB=1.5",lw=0.9),fontsize=8.0)
    ax2.text(0.7,6.25,"A flaperon-scale object is only a few\npixels long; multispectral sampling is coarser\nbefore pan-sharpening.",fontsize=7.2,va="top",linespacing=1.10)
    ax2.text(0.7,4.82,"Possible comparisons",fontsize=8.6,weight="semibold",va="top")
    ax2.text(0.7,4.38,"• elongation and orientation\n• bright area and wake/halo geometry\n• four-band local contrast\n• spatial clustering",fontsize=7.5,linespacing=1.20,va="top")
    ax2.text(0.7,0.45,"Not supportable from report thumbnails alone:\nforensic identification of a specific B777 part.",fontsize=7.5,color="#8A1C2D",weight="semibold",va="bottom",bbox=dict(facecolor="#F8ECEE",edgecolor="#B23A48",pad=3))
    fig.savefig(path,dpi=300,facecolor="white")
    plt.close(fig)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_alt_text(inline_shape, title: str, description: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rPr.append(color); rPr.append(underline); run.append(rPr)
    t=OxmlElement("w:t"); t.text=text; run.append(t); hyperlink.append(run); paragraph._p.append(hyperlink)


def setup_doc(doc: Document):
    sec=doc.sections[0]
    sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    styles=doc.styles
    normal=styles["Normal"]
    normal.font.name="Times New Roman"; normal.font.size=Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
    normal.paragraph_format.line_spacing=1.0
    normal.paragraph_format.space_before=Pt(0); normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for name,size in [("Title",16),("Heading 1",12),("Heading 2",12),("Heading 3",12)]:
        st=styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(size); st.font.bold=False
        st._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
        st.paragraph_format.space_before=Pt(8 if name!="Title" else 0); st.paragraph_format.space_after=Pt(4)
        st.paragraph_format.keep_with_next=True
    styles["Title"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=styles["Caption"]; cap.font.name="Times New Roman"; cap.font.size=Pt(10); cap.font.italic=False
    cap._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
    cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT; cap.paragraph_format.space_after=Pt(5)
    for section in doc.sections:
        footer=section.footer.paragraphs[0]
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=footer.add_run("WORKING DRAFT — 12 August 2026")
        run.font.name="Times New Roman"; run.font.size=Pt(8); run.font.color.rgb=RGBColor(100,100,100)


def add_para(doc, text, style=None, align=None):
    p=doc.add_paragraph(style=style)
    p.add_run(text)
    if align is not None: p.alignment=align
    return p


def add_figure(doc, image_path: Path, width: float, caption: str, title: str, alt: str):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    shape=p.add_run().add_picture(str(image_path),width=Inches(width))
    set_alt_text(shape,title,alt)
    cp=doc.add_paragraph(style="Caption"); cp.paragraph_format.keep_together=True; cp.add_run(caption)
    return cp


SECTION1 = [
"Malaysia Airlines Flight MH370, a Boeing 777-200ER registered as 9M-MRO, was lost on 8 March 2014 with 227 passengers and 12 crew on board. More than twelve years later, the main wreckage and flight recorders have not been located. Several items of debris have been confirmed, or assessed as almost certainly, to have originated from the aircraft, but they do not establish the point or manner of impact. The official safety investigation was consequently unable to determine the cause of the disappearance (Malaysian ICAO Annex 13 Safety Investigation Team, 2018). The unresolved location is therefore not only a search problem. It also prevents the physical examination required to test competing explanations of the accident.",
"The evidential position after the last military-radar contact is unusual. No conventional position reports were received, but the Inmarsat network recorded a small number of signalling exchanges between the aircraft earth station and the Perth ground earth station. The burst timing offset (BTO) constrains the aircraft-to-satellite range, while the burst frequency offset (BFO) contains information about relative velocity after accounting for the satellite motion, the aircraft terminal’s Doppler compensation and frequency biases (Ashton et al., 2015). The received signal level also contains information because the aircraft antenna gain varies with the line-of-sight geometry. None of these is a direct position observation. Each becomes informative only through a physical and statistical model.",
"Further information accumulated after the flight. Aircraft performance and fuel restrict the trajectories that can join the radar track to the satellite-derived range arcs. Debris arrival locations, buoyancy, windage and ocean circulation provide a likelihood for an earlier impact location. Four Pleiades 1A images acquired on 23 March 2014 contain objects assessed as possibly or probably man-made, although the imagery does not establish that they came from MH370 (Minchin et al., 2017). Finally, the areas searched without locating a debris field supply disconfirming, but not logically absolute, evidence: detection probability is high within well-surveyed areas, not equal to one. These information sources refer to different latent states and times. In particular, the aircraft position at the sixth BTO arc at 00:11 UTC, its state during the partial log-on at 00:19 UTC, and its eventual impact point are not interchangeable.",
"Early analyses by Inmarsat and the Australian-led flight path reconstruction group established the southern Indian Ocean as the relevant region. Davey et al. (2016) subsequently formulated the problem as a Bayesian particle filter, combining radar, BTO, BFO, aircraft dynamics and an end-of-flight transition model. Their result was appropriately expressed as a probability density function (PDF), rather than as a single track. They also observed that the assumed Mach-number spread and end-of-flight model materially affected the resulting PDF. This matters because the principal underwater search eliminated a large part of the high-probability region produced under those assumptions, yet did not locate the aircraft (Australian Transport Safety Bureau, 2017). Non-detection does not prove that the earlier estimator was wrong, but it is evidence that should be returned to the estimator rather than considered only after a map has been drawn.",
"The present study therefore treats the problem as one integrated inverse estimate. The baseline inference begins with an uncertain state near the 18:11 UTC radar position and uses BTO, BFO, received power and an aircraft performance/trajectory model through 00:11 UTC. It does not assume that the aircraft was uncontrolled, or that it remained at high altitude, and it does not use the two disputed 00:19 BFO values to determine the sixth-arc airborne state. Alternative controlled and uncontrolled terminal trajectories are then represented by an explicit transition kernel from the airborne state to an impact state. Ocean drift evidence is evaluated at the impact state and marginalised over the unknown terminal trajectory. Search non-detection is likewise introduced as a probabilistic likelihood over the surveyed footprint.",
"The Pleiades hypothesis is handled separately. There is presently no defensible frequency or base-rate estimate from which to assign a probability that the reported objects were MH370 debris. We therefore report two conditional products: one based on the satellite-communications, received-power, performance, drift and search evidence, and a second conditional on the proposition that the relevant Pleiades objects originated from the aircraft. This avoids hiding an essentially judgemental provenance probability within an otherwise numerical posterior.",
"Figure 1 summarises the information flow. The distinction between an observation and a scenario is deliberate. BTO, BFO and received power enter as measurement likelihoods. Aircraft control family, terminal descent or glide, Pleiades provenance and detection performance enter as explicit model states or conditions that can be varied. The paper first establishes the scale and origin of disagreement in prior geospatial estimates, then develops and validates the integrated estimator, and finally reports the posterior maps and their sensitivity to the principal unresolved assumptions.",
]


SECTION2 = [
"2.1. Purpose and scope of the synthesis",
"Large (2019) conducted a meta-analysis of 38 published geospatial estimates: 26 based principally on satellite communications and 12 based on ocean drift. The synthesis remains useful because it shows the scale of between-study dispersion and identifies assumptions that require explicit treatment. It is not, however, used here as a substitute for the observation-level likelihood. Many studies reuse the same Inmarsat log, radar endpoint, aircraft performance information or debris observations and are therefore statistically dependent. Multiplying their published PDFs would count common evidence more than once; averaging their preferred latitudes would discard the physical meaning of the underlying observations.",
"The dissertation converted published point estimates and latitude ranges to a common arc-latitude representation. Where a study did not report a mean or uncertainty, these quantities were reconstructed or imputed, including an assumed Gaussian form for some ranges. The resulting intervals are consequently useful descriptions of the literature, but they should not be read as uniformly calibrated confidence intervals. Figure 2 reproduces the forest plots in a compact form and distinguishes the arithmetic centre of the reported study estimates from the model-based pooled value.",
"2.2. Satellite-communications estimates",
"The 26 satellite-communications estimates had an arithmetic mean of 32.23°S, a sample standard deviation of 7.38° and a range from 12.00°S to 40.55°S. The unweighted fixed-effect centre was 32.23°S (95% interval 31.56–32.90°S), while the weighted random-effects result was 33.34°S (29.13–37.54°S). The heterogeneity statistic was Q = 6368.5 with 25 degrees of freedom (p < 0.0001). The extreme value of Q is more important than the apparent precision of either pooled centre: the studies do not behave as repeated estimates of one location differing only by sampling noise. Their spread reflects materially different trajectory families, speed and turn assumptions, use of BFO, fuel and wind treatments, and end-of-flight kernels, as well as differences in the stated or imputed uncertainty.",
"This result also explains why a consensus latitude is not an adequate prior for a new estimator. A pooled centre would mix different physical models and conceal multi-modality. The present work instead returns to the common measurements, introduces uncertain biases as nuisance states, and compares trajectory and control families within one likelihood framework. The published estimates are retained as an external calibration check: a posterior feature is expected to be explainable by identifiable measurements and assumptions, rather than merely by its proximity to an earlier preferred latitude.",
"2.3. Ocean-drift estimates",
"The 12 drift-study centres were less dispersed: their arithmetic mean was 28.75°S, their sample standard deviation was 2.97° and their reported centres ranged from 23.50°S to 32.00°S. The dissertation’s inverse-variance descriptive pool was 30.17°S (28.74–31.60°S), with Q = 6.55 on 11 degrees of freedom (p = 0.83). Individual reconstructed intervals were much wider, extending in aggregate from approximately 12°S to 38°S. Thus the study centres were comparatively consistent, but the location information supplied by any one drift model remained broad.",
"That apparent agreement must also be interpreted cautiously. The drift studies share ocean reanalyses, debris records and assumptions about leeway, windage and object buoyancy; several estimates are not independent, and reverse drift is intrinsically dispersive. In the integrated estimator, drift is therefore applied as a likelihood for the unknown impact point and then marginalised over terminal-flight displacement. It is not applied directly to the aircraft’s 00:11 position. Alternative circulation products and debris classes are retained as model variants so that agreement between drift models is not mistaken for certainty in their common inputs.",
"2.4. Reproduction strategy",
"The reproduction proceeds at three levels. First, the published BTO arc geometry, BFO residual behaviour and principal aircraft-performance envelopes are reconstructed from the released source data. Secondly, those components are checked against flights and phases for which the aircraft position is known, including MH371 and the known portion of MH370, so that fast random error can be separated from slower or intermittent bias. Thirdly, the components are assembled in a sequential Monte Carlo estimator through 00:11 UTC and compared with the broad features of the Davey et al. (2016) result. Where a proprietary data product or an unpublished implementation prevents exact replication, the substituted model and its sensitivity range are reported explicitly.",
"This staged approach also resolves a potential inconsistency in the literature. Davey et al. (2016) treated the BFOs during the 18:25 and 00:19 transient log-ons as unreliable for the flight-path filter, whereas Holland (2018) later used the last two BFOs in a conditional end-of-flight analysis after bounding oscillator warm-up effects. In the present paper these are not competing choices: the 00:19 BFOs are excluded from the baseline sixth-arc estimate and examined separately as end-of-flight evidence. The resulting posterior through 00:11 is therefore not narrowed by assuming the conclusion of a particular terminal-flight model.",
]


def build_manuscript(path: Path, flow_path: Path, forest_path: Path):
    doc=Document(); setup_doc(doc)
    p=doc.add_paragraph(style="Title"); p.add_run("Flight MH370 Revisited: Integrated Bayesian Estimation without a Prespecified End-of-Flight Scenario")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Large, P. O., Large, J. T., Sarbin, M., ________, ________")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("[Affiliation(s)]")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("[Corresponding author email and ORCID]")
    doc.add_paragraph()
    add_para(doc,"Abstract",style="Heading 1")
    add_para(doc,"[To be drafted after the integrated-model results are fixed; Journal of Navigation limit approximately 150 words.]")
    add_para(doc,"1. INTRODUCTION",style="Heading 1")
    for t in SECTION1: add_para(doc,t)
    add_figure(doc,flow_path,5.85,
        "Figure 1. Structure of the integrated estimator. The baseline state estimate ends at the sixth arc at 00:11 UTC. End-of-flight displacement, ocean drift, Pleiades provenance and search non-detection are represented as explicit transition or likelihood terms rather than being imposed on the airborne state.",
        "Integrated estimator evidence flow",
        "Flow diagram from radar, satellite measurements and aircraft performance to the airborne state at 00:11, then through a marginal terminal-flight kernel to impact, drift, optional Pleiades evidence and search non-detection. The disputed 00:19 observations are reserved for sensitivity analysis.")
    add_para(doc,"2. META-ANALYSIS AND REPRODUCTION",style="Heading 1")
    for t in SECTION2:
        if re.match(r"2\.\d\.",t): add_para(doc,t,style="Heading 2")
        else: add_para(doc,t)
    add_figure(doc,forest_path,5.85,
        "Figure 2. Forest plots of geospatial estimates synthesised in Large (2019): (a) satellite-communications studies and (b) ocean-drift studies. Squares denote study centres and horizontal lines the reported or reconstructed 95% intervals. The diamonds show the dissertation’s random-effects SATCOM result and fixed-effect descriptive drift result. Several intervals were imputed; the estimates share source data and are not independent likelihoods.",
        "Forest plots of prior MH370 geospatial estimates",
        "Two forest plots. Satellite-communications studies span approximately 12 to 41 degrees south and show extreme heterogeneity. Ocean-drift study centres cluster between about 24 and 32 degrees south, although their uncertainty intervals are broad.")
    add_para(doc,"3. INTEGRATED OBSERVATION AND TRAJECTORY MODEL",style="Heading 1")
    add_para(doc,"[Next drafting pass: BTO, BFO, received-power/antenna-gain, aircraft performance, control-family priors and sequential Monte Carlo implementation.]")
    add_para(doc,"4. CALIBRATION, REPRODUCTION AND VALIDATION",style="Heading 1")
    add_para(doc,"[Next drafting pass: known-position checks, MH371/MH370 gain validation, refined BFO noise model, posterior predictive checks and implementation fidelity.]")
    add_para(doc,"5. RESULTS AND SENSITIVITY",style="Heading 1")
    add_para(doc,"[Next drafting pass: 00:11 airborne posterior; drift-marginalised impact posterior; searched-area non-detection; no-power-precompensation sensitivity.]")
    add_para(doc,"6. END-OF-FLIGHT AND CONDITIONAL PLEIADES SCENARIOS",style="Heading 1")
    add_para(doc,"[Next drafting pass: 00:19 evidence, controlled/uncontrolled terminal kernels and a separate posterior conditional on Pleiades provenance.]")
    add_para(doc,"7. CONCLUSIONS AND RECOMMENDATIONS",style="Heading 1")
    add_para(doc,"[To be drafted after the integrated-model and sensitivity results are fixed.]")
    add_para(doc,"REFERENCES",style="Heading 1")
    refs=[
        "Ashton, C., Shuster Bruce, A., Colledge, G. and Dickinson, M. (2015). The Search for MH370. The Journal of Navigation, 68, 1–22. https://doi.org/10.1017/S037346331400068X.",
        "Australian Transport Safety Bureau. (2017). The Operational Search for MH370. External Aviation Investigation AE-2014-054. Canberra: ATSB.",
        "Davey, S., Gordon, N., Holland, I., Rutten, M. and Williams, J. (2016). Bayesian Methods in the Search for MH370. Singapore: SpringerOpen.",
        "Holland, I. D. (2018). MH370 Burst Frequency Offset Analysis and Implications on Descent Rate at End-of-Flight. IEEE Aerospace and Electronic Systems Magazine, 33(2), 24–33. https://doi.org/10.1109/MAES.2018.170048.",
        "Large, P. O. (2019). A Meta-Analysis of Geospatial Estimates in the Case of Malaysian Airlines Flight MH370. Doctoral dissertation, Oklahoma State University.",
        "Malaysian ICAO Annex 13 Safety Investigation Team for MH370. (2018). Safety Investigation Report MH370/01/2018. Putrajaya: Ministry of Transport Malaysia.",
        "Minchin, S., Mueller, N., Lewis, A., Byrne, G. and Tran, M. (2017). Summary of Imagery Analyses for Non-Natural Objects in Support of the Search for Flight MH370. Geoscience Australia Record 2017/13. https://doi.org/10.11636/Record.2017.013.",
    ]
    for r in refs: add_para(doc,r)
    doc.save(path)
    text=' '.join(SECTION1+SECTION2)
    return {"section1":sum(word_count(x) for x in SECTION1),"section2":sum(word_count(x) for x in SECTION2),"combined":word_count(text)}


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"
    hdr=table.rows[0]; set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        cell=hdr.cells[i]; cell.text=h; set_cell_shading(cell,"D9E7F3"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
        for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.size=Pt(font_size)
    for row in rows:
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs: run.font.size=Pt(font_size)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return table


def build_memo(path: Path, holland_path: Path, radar_path: Path, pleiades_path: Path, counts):
    doc=Document(); setup_doc(doc)
    add_para(doc,"MH370 PAPER DEVELOPMENT NOTE",style="Title")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Sections 1–2 budget, end-of-flight evidence, Penang/Pulau Perak altitude sensitivity, and Pleiades image feasibility")
    add_para(doc,"1. EXECUTIVE FINDINGS",style="Heading 1")
    findings=[
        "The Journal of Navigation currently describes research papers as averaging 6,000–8,000 words and up to 20 submission pages including figures and tables. A practical target for this paper is 7,200–7,600 words and 18–20 pages. The first-pass Introduction is %d words and Section 2 is %d words (%d combined), leaving most of the budget for the integrated model, calibration, results and sensitivity analysis."%(counts['section1'],counts['section2'],counts['combined']),
        "Holland’s calculation is strong evidence for rapid descent at the two 00:19 transmission instants if the recorded BFOs are treated as motion-bearing observations after the stated transient and bias corrections. It does not, by itself, identify whether the aircraft was uncontrolled, how long the descent continued, or how far the impact lay from the seventh arc.",
        "The military-radar height read-outs around Penang and Pulau Perak are not reliable measured altitudes. The Malaysian safety report says their variations exceeded aircraft capability and that only latitude and longitude were reasonably reliable. The reported loss of the blip at 18:03:09 and reacquisition at 18:15:25 are nevertheless separate facts about the detection record and should not be discarded with the height labels. The Lido Hotel display was shown in Beijing, not Kuala Lumpur, and its precise provenance and processing remain unpublished.",
        "A Pulau Perak ‘eyewitness’ label can be found in discussion of a leaked police-report appendix, but I have not located the authenticated statement behind it. Later reports confirm that Malaysian Army personnel have been deployed on Pulau Perak, making an observer on military watch plausible. They do not establish staffing on 8 March 2014 or the content and timing of the alleged observation. The combination of gap, reacquisition and witness report merits a joint low-altitude sensitivity analysis, but not yet a fixed low-altitude datum.",
        "All 70 object close-ups and object tables in the Geoscience Australia Pleiades report can be extracted and inspected. Nine of the twelve ‘probably man-made’ objects lie in the north-east PHR_4 scene. The 0.5-m sampling and the report thumbnails are insufficient for a forensic flaperon identification; a probabilistic size/shape/radiance comparison is possible if the original calibrated imagery and appropriate negative controls can be obtained.",
    ]
    for i,f in enumerate(findings,1): add_para(doc,f"{i}. {f}")

    add_para(doc,"2. JOURNAL LENGTH AND SECTION BUDGET",style="Heading 1")
    add_para(doc,"The current Cambridge University Press instructions give an average length of 6,000–8,000 words, a submission length of up to 20 pages including figures and tables, 12-point Times New Roman, single column and 2.54-cm margins. The abstract should be a single paragraph of approximately 150 words and the submission system permits no more than four keywords. The following budget keeps the paper within that envelope while protecting space for the new integrated estimator.")
    add_table(doc,["Component","Target words","Purpose / page pressure"],[
        ["Abstract","150","Finalise after results"],
        ["1. Introduction","800–950","Problem, contribution and evidence architecture"],
        ["2. Meta-analysis and reproduction","750–900","Descriptive synthesis and reproduction logic"],
        ["3. Integrated model","2,000–2,200","Largest methodological section"],
        ["4. Calibration and validation","850–1,000","Known-position and predictive checks"],
        ["5. Results and sensitivity","1,350–1,550","Maps, searched-area update and key variants"],
        ["6. EOF/Pleiades scenarios","750–900","Conditional analyses; no provenance probability"],
        ["7. Conclusions","400–500","Claims, limitations and search implications"],
        ["Total","about 7,250–7,550","Aim for 7–9 figures/tables and 18–20 pages"],
    ],widths=[2.1,1.1,3.3],font_size=9.2)
    add_para(doc,"The two forest plots are scientifically useful but visually expensive. The compact combined figure in the manuscript is suitable for a first submission; if the final paper reaches the page limit, move the individual study rows to supplementary material and retain only pooled centres, ranges and the heterogeneity statistics in the main text.")

    add_para(doc,"3. END-OF-FLIGHT DEEP DIVE: HOLLAND (2017 PREPRINT; 2018 PUBLICATION)",style="Heading 1")
    add_para(doc,"3.1. What was calculated",style="Heading 2")
    add_para(doc,"Holland analysed the BFOs of 182 Hz and −2 Hz recorded eight seconds apart during the 00:19 partial log-on. For level flight near the relevant part of the seventh arc, the expected BFO was approximately 260–280 Hz. The satellite data unit does not compensate for vertical aircraft velocity, giving an approximate sensitivity of 1.7 Hz per 100 ft min−1 of descent. The calculation bounded track, speed, historical BFO error and oven-controlled crystal oscillator (OCXO) warm-up drift under two hypotheses: a fuel-related power interruption followed by auxiliary-power-unit restart, and a short interruption with negligible warm-up drift.")
    add_table(doc,["Holland case","00:19:29 descent","00:19:37 descent","Interpretive status"],[
        ["Fuel/power interruption with bounded warm-up","3,900–14,800 ft min−1","14,800–25,300 ft min−1","Conditional on extrapolated OCXO correction"],
        ["Other/short interruption; no material warm-up","2,900–6,800 ft min−1","13,800–17,600 ft min−1","Conditional on treating recorded BFO normally"],
        ["Combined outer bounds","2,900–14,800 ft min−1","13,800–25,300 ft min−1","Not a joint trajectory posterior"],
    ],widths=[2.0,1.35,1.35,2.0],font_size=9.0)
    add_figure(doc,holland_path,6.45,"Figure A. Evidential layers in the interpretation of the final two BFOs. The measurement-to-descent calculation and the descent-to-impact scenario should be tested separately.","Holland end-of-flight inference layers","Diagram distinguishing recorded BFO facts, signal-model inference of descent, and the further dynamic interpretation as an uncontrolled descent close to the seventh arc, with four principal caveats.")

    add_para(doc,"3.2. Evidence supporting Holland’s narrow conclusion",style="Heading 2")
    for t in [
        "The two recorded values are far below the approximately 260-Hz continuation of the preceding BFO trend. Track changes and ordinary fast random error are much too small, by themselves, to explain the second value.",
        "The vertical-velocity contribution follows from the line-of-sight Doppler geometry and the terminal’s compensation logic; its sign and approximate scale are not controversial.",
        "Earlier 9M-MRO log-ons show a decaying OCXO transient of the sign required to raise the recorded BFO above its steady-state equivalent. Accounting for that effect generally increases, rather than removes, the descent inferred from the recorded values.",
        "Boeing engineering simulations produced uncontrolled cases with descent rates and increases over eight seconds of comparable magnitude, showing that the inferred rates are dynamically possible in at least part of that scenario family.",
    ]: add_para(doc,"• "+t)

    add_para(doc,"3.3. Reasons not to promote the result to a unique end-of-flight model",style="Heading 2")
    for t in [
        "Transient-message status. Davey et al. discarded the 18:25 and 00:19 BFOs from the flight-path filter because the equipment was not in steady state. The last two messages are also of different channel types (R600 and anomalous R1200). Holland’s later treatment is a conditional salvage of those values, not a reversal of their transient character.",
        "Warm-up extrapolation. The seven comparison log-ons followed outages of approximately 20 minutes to more than seven hours; the hypothesised 00:19 outage was approximately one minute. Holland expressly notes that a short outage may produce less drift. The wide correction bounds are sensible for bounding, but not a calibrated transient distribution for a posterior.",
        "Error decomposition. Holland used the historical 4.3-Hz BFO residual standard deviation and outer observed residual bounds. Our short-interval estimate gives about 1.0 Hz for the fast random component once local mean changes are removed. The remaining slow/intermittent bias and start-up transient should be represented as separate latent processes, not folded into one Gaussian error.",
        "Acceleration statistic. The reported approximately 0.68g uses midpoints of two separately bounded intervals. The low bound at the first time and high bound at the second can depend on different track and transient assumptions; the midpoint difference is illustrative, not a posterior estimate of one aircraft trajectory.",
        "Control and continuation. Two line-of-sight velocity observations do not identify attitude, control inputs, altitude, a later recovery, or a glide. Boeing’s close-to-arc result was generated from specified uncontrolled cases, and the ATSB states that the scenarios did not represent every possible condition. Its ±40-NM width was explicitly an envelope for the uncontrolled simulations.",
        "Debris configuration. Damage alignment between the recovered flaperon and outboard flap reduced the likelihood that the main flaps were extended at separation. This weighs against a conventional landing configuration, but it is not a direct observation that nobody was controlling the aircraft.",
    ]: add_para(doc,"• "+t)
    add_para(doc,"A 2024 Journal of Navigation paper by Lyne proposes a controlled eastward descent and directly challenges Holland. It is relevant as a competing published argument, but its stronger claims rely on assumed tracks, a proposed ‘Penang longitude’ destination, pilot-simulator ‘riddles’, JORN avoidance and contested debris interpretation. Its horizontal-speed objection also assumes a particular path between arcs and does not establish that fuel exhaustion before 00:19 is physically impossible. The appropriate response is therefore not to substitute Lyne’s scenario for Holland’s, but to evaluate the final two transmissions under a vector Doppler model while marginalising over track, transient, control and continuation states.")

    add_para(doc,"4. PENANG, THE LIDO DISPLAY AND A LOW-ALTITUDE PULAU PERAK STRESS TEST",style="Heading 1")
    add_para(doc,"4.1. Source assessment",style="Heading 2")
    add_para(doc,"The widely circulated ‘Lido Hotel’ image was photographed during a next-of-kin briefing at the Lido Hotel in Beijing. It was not a Kuala Lumpur display. It resembles a processed or replayed radar presentation, and the raw military returns and processing chain have not been released. The official chronology nevertheless reports that the blip disappeared at 18:03:09, reappeared at 18:15:25 and was then observed until 18:22:12. This interruption is evidence about detection history even though the display is not a calibrated sensor product.")
    add_para(doc,"The 2018 Malaysian safety report is decisive on a narrower point: its height values cannot be read as an altitude profile. It lists blip-to-blip values ranging from approximately 24,450 to 47,500 ft before Penang, approximately 44,700 ft near Penang, 4,800 ft at Pulau Perak, and 29,500 ft after the later reappearance. It also states that the implied altitude and speed changes were beyond aircraft capability, that the extracted values were subject to inherent error, and that only latitude and longitude were reasonably reliable. Figure B therefore separates two propositions that are easily conflated: the height labels are quantitatively unusable, while the reported gap and reacquisition remain observations requiring an explanation.")
    add_figure(doc,radar_path,6.45,"Figure B. Published military-radar height read-outs and the reported contact interruption. Connecting the height values does not create a valid altitude profile, but rejecting those values does not remove the distinct evidence that the blip was lost and later reacquired.","Radar-derived altitude read-outs and contact gap","Line chart showing implausible jumps from roughly 44,700 to 58,200 to 4,800 feet and later 29,500 feet, with the reported loss at 18:03:09 and reacquisition at 18:15:25 marked separately.")
    add_para(doc,"The phrase ‘Pulau Perak – Eyewitness’ appears on a map discussed as page 149 of leaked Police Report Appendix K-2. I have not found a public, authenticated witness statement behind the label. The island is a small, steep and remote rock. Reporting from 2016 and 2019 confirms deployments there by the Malaysian Army’s Fourth Battalion, Royal Ranger Regiment, including duty under Op Pejarak. That makes a military observer, possibly subject to watch routines, a credible class of witness; it does not by itself prove that the post was staffed on the relevant night or establish what any observer saw or heard.")
    add_para(doc,"If the statement can be authenticated as an observation from the island, the evidential combination is more informative than the 4,800-ft label alone. At cruise altitude an airliner’s lights—and, with delay, its sound—can still be detected in clear conditions, so a sighting is not impossible. However, a routine aircraft on a well-used airway would have small apparent size and limited salience to a habituated observer. A substantially lower pass would generally increase angular size, sound level and the probability that the event was noticed and later described as unusual. The size of that likelihood ratio depends on the statement’s actual words, slant range, cloud, moon and background light, aircraft lighting, observer location and duty, and the delay and selection process by which the report entered the police material.")
    add_para(doc,"A simple scale calculation illustrates the contrast without pretending to estimate a witness likelihood. The Boeing 777-200ER span is 60.9 m. For a directly overhead pass it would subtend about 0.33° at 35,000 ft, 0.57° at 20,000 ft, 1.15° at 10,000 ft and 2.29° at 5,000 ft. The corresponding sound-arrival delays are approximately 31, 18, 9 and 4 s. Geometric spreading alone would make an otherwise comparable sound about 11 dB stronger at 10,000 ft and 17 dB stronger at 5,000 ft than at 35,000 ft. These figures do not account for lateral offset, engine thrust, atmospheric absorption, cloud, aircraft lighting, sea and background noise, or observer attention. They do quantify why a low pass would be much more salient to a watchstander while leaving a cruise-altitude observation physically possible.")

    add_para(doc,"4.2. How to test the descent–reclimb hypothesis",style="Heading 2")
    add_para(doc,"The integrated estimator should add an early-altitude scenario variable rather than force the 4,800-ft read-out. A useful family would sample a minimum altitude between 5,000 and 20,000 ft, a low-level duration between zero and ten minutes, and a climb ending between 18:15 and 18:25 at 28,000–35,000 ft. It should be conditioned on the reliable radar latitude/longitude track and on Boeing 777 climb-rate, calibrated-airspeed, cabin-pressurisation and fuel-flow constraints. If the aircraft was still climbing at the 18:25 log-on, the vertical BFO term and antenna geometry must also enter directly.")
    add_para(doc,"The radar and witness terms should then be evaluated jointly, not counted as independent votes. Let G denote the reported gap/reacquisition and W the authenticated contents of the witness account. For each sampled trajectory T, the calculation should marginalise P(G | T, radar site, beam and processing parameters) and P(W | T, visibility, acoustics, observer and reporting parameters). A low trajectory can increase both probabilities, creating dependence through altitude. Competing families should include (i) cruise altitude with a coverage, propagation or processing gap; (ii) descent below effective coverage followed by reacquisition during a climb; and (iii) an incorrect track or witness association. Until raw radar data and the statement are obtained, the result should be reported across explicit likelihood-ratio ranges rather than as a single Bayes factor.")
    add_para(doc,"A first energy/fuel bracket indicates the scale, not a final estimator result. Re-climbing by 25,000–30,000 ft at a mass near 210 tonnes requires approximately 15–19 GJ of additional potential energy. Across a broad 25–40% effective conversion from fuel energy to aircraft potential energy, that corresponds to approximately 0.9–1.8 tonnes of fuel before crediting the lower thrust used during descent. Descent fuel saving and the duration of low-level flight make a net penalty of roughly 0.3–1.3 tonnes a defensible stress range. At a later two-engine cruise flow of about 5.5–7.0 tonnes per hour, this is equivalent to approximately 3–14 minutes of endurance, or roughly 20–110 NM at 430–480 kt.")
    add_para(doc,"The sixth-arc BTO geometry would move only a few nautical miles for the lower altitude itself. The larger effect is reweighting: a fuel penalty removes trajectories with little remaining endurance and tends to favour shorter-range solutions, while a climb still in progress at 18:25 also changes the BFO likelihood. The direction and magnitude of any shift in the two-dimensional PDF cannot be stated from the energy bracket alone. It requires a particle-level rerun with fuel mass carried as a state. The existing working estimator uses a broad endurance/performance envelope, so quoting a precise latitude shift at this stage would be false precision.")

    add_para(doc,"5. PLEIADES IMAGE AND DEBRIS-COMPONENT FEASIBILITY",style="Heading 1")
    add_para(doc,"Geoscience Australia examined four Pleiades 1A scenes acquired on 23 March 2014 at 0.5-m pan-sharpened resolution. Seventy objects were catalogued: 12 assessed as ‘probably man-made’ and 28 as ‘possibly man-made’. The north-east PHR_4 scene contains 36 catalogued objects and nine of the twelve highest-rated candidates. Their reported bright-object footprints range from 23 to 70 m². The automated spectral classifications were not reliable because glint and shaded glint resembled the candidate objects; the final ratings relied substantially on manual inspection and principal-components displays.")
    add_figure(doc,pleiades_path,6.45,"Figure C. Left: positions and reported bright-footprint areas of the nine PHR_4 objects rated ‘probably man-made’. Right: a sampling-scale illustration showing why component identification is under-resolved. This is a new data-derived schematic; it does not reproduce the CNES image pixels.","Pleiades candidate-object map and resolution schematic","Map of nine high-rated objects in the north-east Pleiades scene and a pixel grid showing that a two-metre object spans only about four 0.5-m pixels, followed by feasible and infeasible image-analysis claims.")
    add_table(doc,["Scene / object","Latitude","Longitude","Reported area"],[[f"PHR_4 / {i}",f"{lat:.5f}°",f"{lon:.6f}°",f"{area} m²"] for i,lat,lon,area in PHR4_RATING5],widths=[1.3,1.4,1.5,1.2],font_size=9.0)
    add_para(doc,"I can extract and catalogue every true-colour and principal-components close-up in the report, register the reported coordinates and metrics, and compare candidate morphology against dimensions of the flaperon, outboard-flap sections, fairings, aileron and other confirmed or highly likely debris. The report imagery alone supports only a low-dimensional probabilistic comparison: bounding-box length/width where recoverable, elongation, orientation, bright area, wake or halo geometry, four-band contrast relative to local sea surface, and clustering. A recovered flaperon roughly two metres long would span only about four pixels at 0.5-m ground sampling; radiometric mixing, the sensor point-spread function, pan-sharpening and sea-surface motion prevent a reliable visual identification.")
    add_para(doc,"A stronger analysis would require the original radiometrically calibrated Pleiades products (not screenshots), acquisition metadata and point-spread information; a set of negative-control scenes with comparable sea state; exact component dimensions and flotation attitudes; and a pre-registered scoring rule tested without knowledge of the candidate labels. The report itself notes the absence of comparable negative-control images. Spectral comparison should be described as four-band radiance/contrast analysis, not material spectroscopy: an object only a few pixels across cannot yield a clean material spectrum from these data.")
    add_para(doc,"There is also a rights issue. The Geoscience Australia text is generally CC BY 4.0, but the embedded imagery is marked French Military Intelligence Service © CNES and is expressly excepted from that licence. I can use the close-ups for internal scientific inspection; reproduction in a journal article would require a permission or licensing check. Newly drawn coordinate maps and data tables, such as Figure C, avoid that problem.")

    add_para(doc,"6. RECOMMENDED NEXT TESTS",style="Heading 1")
    tests=[
        "Implement the early descent–reclimb family as a latent initial-condition/fuel-penalty model and rerun the estimator through 00:11, reporting the Bayes factor or posterior mass relative to the level-flight family and the change in the two-dimensional sixth-arc PDF.",
        "Build a dedicated 00:19 likelihood with separate fast BFO noise, slow/intermittent bias, channel offset and OCXO-transient states; compare controlled continuation, recovery/glide and uncontrolled simulation families without making any one of them the baseline.",
        "Acquire or request the underlying Pulau Perak witness statement and 8 March 2014 post roster/watch information before assigning it a likelihood. Until then, retain the gap-plus-witness combination as a labelled joint stress condition with broad radar-detection and human-observation models.",
        "Extract the 70 Pleiades object records into a reproducible catalogue, acquire original imagery if possible, and construct negative-control object detections before attempting any flaperon or debris-class score.",
        "Keep the Pleiades output conditional: publish one integrated PDF without that hypothesis and one conditional on it, without inventing a prior probability that the objects came from MH370.",
    ]
    for i,t in enumerate(tests,1): add_para(doc,f"{i}. {t}")

    add_para(doc,"PRIMARY SOURCES CONSULTED",style="Heading 1")
    sources=[
        ("Journal of Navigation author instructions","https://www.cambridge.org/core/journals/journal-of-navigation/information/author-instructions/preparing-your-materials"),
        ("Holland preprint and publication record","https://arxiv.org/abs/1702.02432"),
        ("ATSB MH370 investigation and end-of-flight material","https://www.atsb.gov.au/investigations/ae-2014-054"),
        ("Malaysian Ministry of Transport MH370 archive","https://www.mot.gov.my/en/aviation/reports/archived-report/mh370/"),
        ("Later reporting confirming Malaysian Army deployments on Pulau Perak","https://www.astroawani.com/berita-malaysia/mayat-koperal-tentera-dijumpai-dekat-pulau-perak-118578"),
        ("Boeing 777-200ER dimensions","https://www.boeing.com/commercial/777"),
        ("ATSB / Geoscience Australia Pleiades report page","https://www.atsb.gov.au/summary-imagery-analyses-non-natural-objects-support-search-flight-mh370"),
        ("Lyne (2024), Journal of Navigation","https://doi.org/10.1017/S0373463324000262"),
    ]
    for label,url in sources:
        p=doc.add_paragraph(); add_hyperlink(p,label,url)
    doc.save(path)


def main():
    flow=FIG/"Figure_1_integrated_estimator_flow.png"
    forest=FIG/"Figure_2_meta_analysis_forest_plots.png"
    holland=FIG/"Memo_Figure_A_Holland_inference_layers.png"
    radar=FIG/"Memo_Figure_B_radar_altitude_readouts.png"
    pleiades=FIG/"Memo_Figure_C_Pleiades_candidates_resolution.png"
    make_evidence_flow(flow); make_forest(forest); make_holland_figure(holland); make_radar_figure(radar); make_pleiades_figure(pleiades)
    manuscript=OUT/"Flight_MH370_Revisited_Sections_1_2_First_Pass.docx"
    memo=OUT/"MH370_End_of_Flight_Radar_and_Pleiades_Evidence_Note.docx"
    counts=build_manuscript(manuscript,flow,forest)
    build_memo(memo,holland,radar,pleiades,counts)
    (OUT/"word_count_and_budget.txt").write_text(
        f"Section 1: {counts['section1']} words\nSection 2: {counts['section2']} words\nCombined Sections 1-2: {counts['combined']} words\nTarget whole paper: 7,200-7,600 words; 18-20 submission pages.\n",
        encoding="utf-8",
    )
    print(manuscript); print(memo); print(counts)


if __name__ == "__main__":
    main()
