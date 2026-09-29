#!/usr/bin/env python3
"""Build the inputs of the MH371 known-flight control (9M-MRO, 7 March 2014, Beijing to Kuala Lumpur).

Usage: .venv/bin/python .sources/mh371/build.py SITA_XLSX SATELLITE_XLSX ACARS_XLSX

Writes data/mh371/satcom-observations.csv and data/mh371/satellite-ephemeris.csv, and prints
the [prior] block for config/mh371.toml. The control is blind: everything below was fixed
before any run, and the truth track is read only by report/mh371_score.py.

Window: Davey et al. (2016) Fig. 9.7, the published validation segment, 01:59 to 06:54 UTC
(read from the figure's dotted lines to about a minute).

Measurements, chosen by rule from the SITA SU log (not by looking at the flight):
- R1200 bursts with a BTO, received through Perth (GES 305) on the Indian Ocean Region
  satellite (Inmarsat-3F1, the satellite of data/satellite-ephemeris.csv), CRC correct.
- Excluding log-on/log-off units, whose BTOs carry the anomalous offsets (Davey sec. 5.2).
- One burst per hour, MH370's cadence: the first qualifying burst at or after T0 + k h
  (k = 1, 2, ...), and the last qualifying burst before the window end as the final epoch.
- BTO only. The satellite + EAFC frequency terms that the BFO model needs are tabulated
  (ATSB) only for MH370's bursts, so the BFO is not used (`cruise` = bto). The
  satellite_afc_hz column is 0 and unused.

Prior (the only read of ACARS here, declared): the last ACARS report at or before the window
start gives the prior time, position and true heading. The heading stands in for the track
with s.d. 5 deg for the unknown wind crab (MH370's radar track has 1 deg). An earlier version
took the bearing between the last two reports; it was replaced before any filter run because
that pair spans the turn near 01:58 shown in Davey's Fig. 9.7. Mach and altitude keep the
base prior (uniform), as for MH370; the position s.d. is MH370's 0.5 NM.
"""

import csv
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data/mh371"
DAY = datetime(2014, 3, 7, tzinfo=timezone.utc)
WINDOW = ("01:59:00", "06:54:00")
TYPES_EXCLUDED = ("Log-on", "Log-off", "Log Control")
BTO_SD_US = 29.0


def seconds(hms):
    h, m, s = hms.split(":")
    return int(h) * 3600 + int(m) * 60 + float(s)


def utc(sec_of_day):
    t = DAY.timestamp() + sec_of_day
    return datetime.fromtimestamp(t, timezone.utc)


def acars_prior(path, t0):
    """Time, position and true heading (column HDG(T)) of the last ACARS report at or before t0."""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)["ACARS"].iter_rows(values_only=True))
    reports = [(r[11], r[17], r[18], r[23]) for r in rows
               if len(r) > 23 and all(isinstance(r[k], (int, float)) for k in (11, 17, 18, 23))]
    return max(r for r in reports if r[0] <= t0)


def main():
    sita, satellite, acars = map(Path, sys.argv[1:4])
    start, end = (seconds(w) for w in WINDOW)
    t0, lat0, lon0, track0 = acars_prior(acars, start)

    rows = list(openpyxl.load_workbook(sita, read_only=True, data_only=True)["SU Log"].iter_rows(values_only=True))
    bursts = []
    for i, r in enumerate(rows[2:], start=3):
        if not r[0] or not str(r[0]).startswith("2014-03-07"):
            continue
        t = seconds(str(r[0])[11:])
        if not (t0 < t <= end) or r[27] in (None, ""):
            continue
        if "-R1200-" not in str(r[3]) or r[4] != "IOR" or str(r[5]) != "305" or r[20] != "Yes":
            continue
        if any(x in str(r[13]) for x in TYPES_EXCLUDED):
            continue
        bursts.append((t, float(r[27]), i, str(r[13])))
    bursts.sort()
    chosen, k = [], 1
    while True:
        after = [b for b in bursts if b[0] >= t0 + 3600 * k]
        if not after:
            break
        chosen.append(after[0])
        k += 1
    if bursts[-1] not in chosen:
        chosen.append(bursts[-1])

    sat = {}
    for r in openpyxl.load_workbook(satellite, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True):
        if isinstance(r[1], (int, float)) and isinstance(r[2], (int, float)):
            sat[int(r[1])] = r[2:8]

    OUT.mkdir(parents=True, exist_ok=True)
    with open(OUT / "satcom-observations.csv", "w", newline="") as f, open(OUT / "satellite-ephemeris.csv", "w", newline="") as g:
        obs, eph = csv.writer(f), csv.writer(g)
        obs.writerow(["epoch_id", "time_utc", "seconds_from_t0", "bto_us", "bto_sd_us", "bfo_hz", "bfo_sd_hz",
                      "satellite_afc_hz", "raw_source_rows", "logged_utc", "cruise", "events", "note"])
        eph.writerow(["epoch_id", "time_utc", "x_km", "y_km", "z_km", "vx_km_s", "vy_km_s", "vz_km_s"])
        for t, bto, row, kind in chosen:
            whole = int(t)
            stamp = utc(whole).strftime("%Y-%m-%dT%H:%M:%S.000Z")
            eid = "k" + utc(whole).strftime("%H%M")
            obs.writerow([eid, stamp, f"{whole - t0:.1f}", f"{bto:.0f}", BTO_SD_US, "", "", "0",
                          row, utc(t).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z", "bto", "", f"R1200 {kind}; BTO only"])
            x, y, z, vx, vy, vz = sat[whole]
            eph.writerow([eid, stamp, x, y, z, vx, vy, vz])
            print(f"{eid} {stamp} BTO {bto:.0f} ({kind}, SU log row {row})")
    print(f"\nwrote {OUT}/satcom-observations.csv and satellite-ephemeris.csv ({len(chosen)} epochs)")
    print(f'\n[prior]\ntime_utc = "{utc(t0).strftime("%Y-%m-%dT%H:%M:%SZ")}"\nlatitude_deg = {lat0:.6f}\n'
          f"longitude_deg = {lon0:.6f}\nposition_sd_nm = 0.5\ntrack_deg = {track0 % 360:.2f}\ntrack_sd_deg = 5.0")


if __name__ == "__main__":
    main()
