#!/usr/bin/env python3
"""Extract a compact, post-selection ACARS track overlay.

The 31 MB source workbook is distributed separately in the user-provided
dissertation-spreadsheet archive and is deliberately not copied into this
source recreation.  This adapter refuses any workbook other than the audited
file, selects only the non-duplicate position rows in its ``ACARS`` sheet, and
writes a deterministic CSV plus a provenance manifest.

The output is descriptive map context.  It is not an input to WSPR anomaly
selection, PHaRLAP screening, clustering, residualization, or truth scoring.
Consecutive reported positions may be joined for display, but the joins must
not be described as continuously observed aircraft positions.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import platform
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Mapping, Sequence


SOURCE_FILENAME = "ACARS POSITION AND WIND REPORTS.xlsx"
SOURCE_SHA256 = "5c9a79d106fbad3900db36b5ddf9d3baa57b8de5a76e72b95fff301fee6d81fe"
SOURCE_SIZE_BYTES = 31_260_031
SOURCE_URL = (
    "https://huggingface.co/datasets/peteabiome/"
    "mh370-dissertation-spreadsheets/resolve/main/"
    "ACARS%20POSITION%20AND%20WIND%20REPORTS.xlsx?download=true"
)
DATASET_URL = (
    "https://huggingface.co/datasets/peteabiome/"
    "mh370-dissertation-spreadsheets"
)
SOURCE_SHEET = "ACARS"
OBSERVATION_DATE = (2014, 3, 7)
TRACK_SCHEMA = "mh370.wspr.known-flight-track-observations.v1"
MANIFEST_SCHEMA = "mh370.wspr.known-flight-track-provenance.v1"

# These ranges exclude repeated copies at rows 98--102 and 116--117.
SOURCE_ROW_RANGES = {
    "MH371": (14, 95),
    "MH370": (105, 110),
}
EXPECTED_COUNTS = {"MH371": 71, "MH370": 6}
DUPLICATE_ROW_MAP = {
    98: 91,
    99: 92,
    100: 93,
    101: 94,
    102: 95,
    116: 14,
    117: 15,
}

# One-based workbook columns.
COLUMNS = {
    "hour": 9,
    "minute": 10,
    "second": 11,
    "altitude_ft": 13,
    "latitude_deg": 18,
    "longitude_deg_e": 19,
    "heading_true_deg": 29,
}
OUTPUT_FIELDS = (
    "flight",
    "track_point_index",
    "time_utc",
    "latitude_deg",
    "longitude_deg_e",
    "altitude_ft",
    "heading_true_deg",
    "position_source",
    "position_method",
    "source_sheet",
    "source_excel_row",
)
NUMERIC_SOURCE_FIELDS = tuple(COLUMNS)


class TrackExtractionError(RuntimeError):
    """The supplied workbook or extracted track violated the frozen audit."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def canonical_json(value: object) -> bytes:
    return (
        json.dumps(value, sort_keys=True, indent=2, allow_nan=False) + "\n"
    ).encode("utf-8")


def csv_bytes(rows: Sequence[Mapping[str, object]]) -> bytes:
    destination = io.StringIO(newline="")
    writer = csv.DictWriter(
        destination, fieldnames=OUTPUT_FIELDS, lineterminator="\n"
    )
    writer.writeheader()
    writer.writerows(rows)
    return destination.getvalue().encode("utf-8")


def atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    except BaseException:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise


def verify_source(path: Path) -> None:
    if not path.is_file():
        raise TrackExtractionError(f"source workbook not found: {path}")
    size = path.stat().st_size
    if size != SOURCE_SIZE_BYTES:
        raise TrackExtractionError(
            f"source workbook size is {size}, expected {SOURCE_SIZE_BYTES}"
        )
    identity = sha256_file(path)
    if identity != SOURCE_SHA256:
        raise TrackExtractionError(
            f"source workbook SHA-256 is {identity}, expected {SOURCE_SHA256}"
        )


def _numeric_values(sheet: object, row_number: int) -> dict[str, float] | None:
    values = {
        field: sheet.cell(row_number, column).value
        for field, column in COLUMNS.items()
    }
    present = [value is not None for value in values.values()]
    if not any(present):
        return None
    if not all(
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(float(value))
        for value in values.values()
    ):
        # Header/separator rows within the selected MH371 range are expected.
        if values["hour"] is None and values["latitude_deg"] is None:
            return None
        raise TrackExtractionError(
            f"{SOURCE_SHEET}!{row_number}: incomplete or non-numeric track row"
        )
    return {field: float(value) for field, value in values.items()}


def _source_signature(sheet: object, row_number: int) -> tuple[float, ...]:
    values = _numeric_values(sheet, row_number)
    if values is None:
        raise TrackExtractionError(
            f"{SOURCE_SHEET}!{row_number}: expected a numeric duplicate row"
        )
    return tuple(values[field] for field in NUMERIC_SOURCE_FIELDS)


def extract_rows(workbook_path: Path) -> tuple[list[dict[str, object]], str]:
    verify_source(workbook_path)
    try:
        import openpyxl
        from openpyxl import load_workbook
    except ImportError as error:
        raise TrackExtractionError("openpyxl is required for extraction") from error

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    if SOURCE_SHEET not in workbook.sheetnames:
        raise TrackExtractionError(f"workbook lacks sheet {SOURCE_SHEET!r}")
    sheet = workbook[SOURCE_SHEET]

    for duplicate, original in DUPLICATE_ROW_MAP.items():
        if _source_signature(sheet, duplicate) != _source_signature(sheet, original):
            raise TrackExtractionError(
                f"{SOURCE_SHEET}!{duplicate} no longer duplicates row {original}"
            )

    rows: list[dict[str, object]] = []
    per_flight: dict[str, list[dict[str, object]]] = {}
    year, month, day = OBSERVATION_DATE
    for flight, (first_row, last_row) in SOURCE_ROW_RANGES.items():
        flight_rows: list[dict[str, object]] = []
        for row_number in range(first_row, last_row + 1):
            values = _numeric_values(sheet, row_number)
            if values is None:
                continue
            hour, minute, second = (
                int(values["hour"]),
                int(values["minute"]),
                int(values["second"]),
            )
            if not (
                values["hour"] == hour
                and values["minute"] == minute
                and values["second"] == second
                and 0 <= hour < 24
                and 0 <= minute < 60
                and 0 <= second < 60
            ):
                raise TrackExtractionError(
                    f"{SOURCE_SHEET}!{row_number}: invalid UTC components"
                )
            latitude = values["latitude_deg"]
            longitude = values["longitude_deg_e"]
            heading = values["heading_true_deg"]
            if not -90.0 <= latitude <= 90.0:
                raise TrackExtractionError(
                    f"{SOURCE_SHEET}!{row_number}: invalid latitude"
                )
            if not -180.0 <= longitude <= 180.0:
                raise TrackExtractionError(
                    f"{SOURCE_SHEET}!{row_number}: invalid longitude"
                )
            if not 0.0 <= heading < 360.0:
                raise TrackExtractionError(
                    f"{SOURCE_SHEET}!{row_number}: invalid true heading"
                )
            timestamp = datetime(
                year, month, day, hour, minute, second, tzinfo=timezone.utc
            )
            flight_rows.append({
                "flight": flight,
                "track_point_index": len(flight_rows),
                "time_utc": timestamp.isoformat(timespec="seconds").replace(
                    "+00:00", "Z"
                ),
                "latitude_deg": latitude,
                "longitude_deg_e": longitude,
                "altitude_ft": values["altitude_ft"],
                "heading_true_deg": heading,
                "position_source": SOURCE_FILENAME,
                "position_method": "reported ACARS position sample",
                "source_sheet": SOURCE_SHEET,
                "source_excel_row": row_number,
            })
        if len(flight_rows) != EXPECTED_COUNTS[flight]:
            raise TrackExtractionError(
                f"{flight}: extracted {len(flight_rows)} rows, "
                f"expected {EXPECTED_COUNTS[flight]}"
            )
        timestamps = [row["time_utc"] for row in flight_rows]
        if timestamps != sorted(set(timestamps)):
            raise TrackExtractionError(f"{flight}: timestamps are not unique and ordered")
        per_flight[flight] = flight_rows
        rows.extend(flight_rows)
    return rows, openpyxl.__version__


def generate(
    workbook_path: Path, output_directory: Path
) -> tuple[Path, Path, dict[str, object]]:
    rows, openpyxl_version = extract_rows(workbook_path)
    csv_payload = csv_bytes(rows)
    csv_name = "known_flight_track_observations.csv"
    manifest_name = "known_flight_track_provenance.json"
    csv_path = output_directory / csv_name
    manifest_path = output_directory / manifest_name
    counts = {
        flight: sum(row["flight"] == flight for row in rows)
        for flight in SOURCE_ROW_RANGES
    }
    time_ranges = {
        flight: {
            "first": next(
                row["time_utc"] for row in rows if row["flight"] == flight
            ),
            "last": next(
                row["time_utc"]
                for row in reversed(rows)
                if row["flight"] == flight
            ),
        }
        for flight in SOURCE_ROW_RANGES
    }
    manifest: dict[str, object] = {
        "schema": MANIFEST_SCHEMA,
        "track_schema": TRACK_SCHEMA,
        "role": (
            "Post-selection descriptive overlay only; never supplied to WSPR "
            "selection, PHaRLAP screening, clustering, residualization or "
            "truth scoring. Lines between rows are display interpolation, not "
            "continuously observed positions."
        ),
        "source": {
            "dataset": "MH370 Dissertation Working Spreadsheets",
            "dataset_url": DATASET_URL,
            "file": SOURCE_FILENAME,
            "download_url": SOURCE_URL,
            "sha256": SOURCE_SHA256,
            "size_bytes": SOURCE_SIZE_BYTES,
            "license": "CC-BY-4.0",
            "worksheet": SOURCE_SHEET,
            "workbook_creator": "Pete Large",
            "workbook_created_utc": "2019-02-18T18:50:06Z",
            "workbook_modified_utc": "2019-02-18T22:32:04Z",
        },
        "selection": {
            "observation_date_utc": "2014-03-07",
            "inclusive_excel_row_ranges": {
                flight: list(bounds)
                for flight, bounds in SOURCE_ROW_RANGES.items()
            },
            "excluded_repeated_excel_rows": sorted(DUPLICATE_ROW_MAP),
            "repeated_row_sources": {
                str(duplicate): original
                for duplicate, original in DUPLICATE_ROW_MAP.items()
            },
            "rule": (
                "Within the declared flight row ranges retain rows whose UTC "
                "components, altitude, latitude, longitude and normalized true "
                "heading cells are all finite numeric values."
            ),
            "workbook_columns_one_based": COLUMNS,
        },
        "output": {
            "file": csv_name,
            "sha256": hashlib.sha256(csv_payload).hexdigest(),
            "rows": len(rows),
            "rows_by_flight": counts,
            "time_ranges_utc": time_ranges,
            "columns": list(OUTPUT_FIELDS),
        },
        "runtime": {
            "python": platform.python_version(),
            "openpyxl": openpyxl_version,
        },
    }
    atomic_write(csv_path, csv_payload)
    atomic_write(manifest_path, canonical_json(manifest))
    return csv_path, manifest_path, manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output-directory",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "data" / "known-flight",
    )
    return parser


def main() -> int:
    arguments = build_parser().parse_args()
    csv_path, manifest_path, manifest = generate(
        arguments.workbook, arguments.output_directory
    )
    print(json.dumps({
        "csv": str(csv_path),
        "manifest": str(manifest_path),
        "rows": manifest["output"]["rows"],
        "sha256": manifest["output"]["sha256"],
    }, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
