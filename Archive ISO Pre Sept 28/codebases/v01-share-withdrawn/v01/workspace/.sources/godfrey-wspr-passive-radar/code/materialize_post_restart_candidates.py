#!/usr/bin/env python3
"""Freeze trajectory-blind WSPR candidates at the seven post-restart BTO arcs.

This is source-recreation code, not estimator product code.  It applies the
published public WSPR anomaly rule to a frozen WSPR Live extract, constructs
the eastern BTO-arc branch from measured SATCOM inputs, and retains every
same-slot anomalous-link intersection in a prespecified broad map domain and
within 100 NM of an arc.  No assumed aircraft trajectory or post-radar truth
position enters candidate construction.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import io
import json
import math
import os
import platform
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.spatial import cKDTree

import wspr_pharlap_control as control


maidenhead_six_character_center = control.maidenhead_six_character_center


SCHEMA = "mh370.wspr.post-restart-candidates.v1"
MAP_EXTENT = (64.0, 110.0, -45.0, 10.0)  # west, east, south, north
BAND_NM = 100.0
WINDOW_MINUTES = 10.0
BASELINE_HOURS = 3.0
MINIMUM_BASELINE_ROWS = 5
ANOMALY_THRESHOLD = 0.75
MINIMUM_CROSSING_ANGLE_DEG = 20.0
WSPR_MIDPOINT_OFFSET_SECONDS = 56.242
CONTROL_OFFSETS_MINUTES = {
    "minus_60_min": -60,
    "actual_time": 0,
    "plus_60_min": 60,
}
CONDITION_ORDER = {name: index for index, name in enumerate(CONTROL_OFFSETS_MINUTES)}
INCLUDED_BANDS = (-1, 0, 1, 3, 5, 7, 10, 14, 18, 21, 24, 28)

C_KM_S = 299_792.458
BTO_BIAS_US = -495_679.0
GES_ECEF_KM = np.array((-2368.8, 4881.1, -3342.0), dtype=float)
ARC_ALTITUDE_KM = 10.0
WGS84_A_KM = 6378.137
WGS84_E2 = 6.69437999014e-3
EARTH_MEAN_RADIUS_KM = 6371.0088
ARC_LATITUDE_GRID_POINTS = 14_001

LOCATOR_PATTERN = re.compile(r"^[A-Ra-r]{2}[0-9]{2}[A-Xa-x]{2}$")
REQUIRED_WSPR_COLUMNS = (
    "id", "time", "band", "rx_sign", "rx_lat", "rx_lon", "rx_loc",
    "tx_sign", "tx_lat", "tx_lon", "tx_loc", "frequency", "snr",
)
SPOT_FIELDS = (
    "id", "time", "frequency", "tx_sign", "rx_sign", "tx_loc", "rx_loc",
    "tx_lat", "tx_lon", "rx_lat", "rx_lon",
)


EPOCHS: dict[str, dict[str, object]] = {
    "m1825": {
        "sequence": 1,
        "title": "Restart / first arc",
        "handshake": "2014-03-07 18:25:34",
        "bto_us": 12_600.0,
        "ephemeris_second": 66_334,
        "bto_provenance": (
            "canonical estimator observation m1825, raw source row 7013; "
            "anomalous R1200, BFO excluded"
        ),
    },
    "m1941": {
        "sequence": 2,
        "title": "Second arc",
        "handshake": "2014-03-07 19:41:02",
        "bto_us": 11_500.0,
        "ephemeris_second": 70_862,
        "bto_provenance": "canonical estimator observation m1941, raw source row 7126",
    },
    "m2041": {
        "sequence": 3,
        "title": "Third arc",
        "handshake": "2014-03-07 20:41:04",
        "bto_us": 11_740.0,
        "ephemeris_second": 74_464,
        "bto_provenance": "canonical estimator observation m2041, raw source row 7128",
    },
    "m2141": {
        "sequence": 4,
        "title": "Fourth arc",
        "handshake": "2014-03-07 21:41:26",
        "bto_us": 12_780.0,
        "ephemeris_second": 78_086,
        "bto_provenance": "canonical estimator observation m2141, raw source row 7130",
    },
    "m2241": {
        "sequence": 5,
        "title": "Fifth arc",
        "handshake": "2014-03-07 22:41:21",
        "bto_us": 14_540.0,
        "ephemeris_second": 81_681,
        "bto_provenance": "canonical estimator observation m2241, raw source row 7132",
    },
    "m0011": {
        "sequence": 6,
        "title": "Sixth arc",
        "handshake": "2014-03-08 00:10:59",
        "bto_us": 18_040.0,
        "ephemeris_second": 87_059,
        "bto_provenance": "canonical estimator observation m0011, raw source row 7187",
    },
    "m0019a": {
        "sequence": 7,
        "title": "Seventh arc",
        "handshake": "2014-03-08 00:19:29",
        "bto_us": 18_400.0,
        "raw_bto_us": 23_000.0,
        "r600_correction_us": 4_600.0,
        "ephemeris_second": 87_569,
        "bto_provenance": (
            "canonical estimator observation m0019a, raw source row 7188; "
            "R600 log-on BTO after declared 4600-us correction, BFO excluded"
        ),
    },
}


class MaterializationError(RuntimeError):
    """A frozen input or scientific invariant is invalid."""


@dataclass(frozen=True)
class Arc:
    epoch_id: str
    handshake: pd.Timestamp
    bto_us: float
    satellite_ecef_km: np.ndarray
    target_range_km: float
    latitude_deg: np.ndarray
    longitude_deg_e: np.ndarray
    unit_vectors: np.ndarray
    index: cKDTree


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def canonical_json(value: object) -> bytes:
    return (json.dumps(value, sort_keys=True, indent=2, allow_nan=False) + "\n").encode()


def deterministic_gzip(payload: bytes) -> bytes:
    destination = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=destination, mtime=0) as handle:
        handle.write(payload)
    return destination.getvalue()


def atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    except BaseException:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise


def dataframe_csv_bytes(frame: pd.DataFrame) -> bytes:
    return frame.to_csv(
        None, index=False, lineterminator="\n", float_format="%.15g"
    ).encode()


def unit_vectors(latitude_deg: Iterable[float], longitude_deg_e: Iterable[float]) -> np.ndarray:
    latitude = np.radians(np.asarray(latitude_deg, dtype=float))
    longitude = np.radians(np.asarray(longitude_deg_e, dtype=float))
    cosine = np.cos(latitude)
    return np.column_stack((
        cosine * np.cos(longitude), cosine * np.sin(longitude), np.sin(latitude)
    ))


def read_ephemeris(path: Path) -> dict[int, np.ndarray]:
    """Read exact one-second satellite ECEF rows from the retained workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    wanted = {int(value["ephemeris_second"]) for value in EPOCHS.values()}
    result: dict[int, np.ndarray] = {}
    for row in sheet.iter_rows(min_row=1, values_only=True):
        second = row[1]
        if second is None:
            continue
        key = int(second)
        if key in wanted:
            result[key] = np.asarray(row[2:5], dtype=float)
            if len(result) == len(wanted):
                break
    workbook.close()
    missing = wanted.difference(result)
    if missing:
        raise MaterializationError(f"ephemeris lacks seconds {sorted(missing)}")
    return result


def construct_arc(epoch_id: str, satellite_ecef_km: np.ndarray) -> Arc:
    case = EPOCHS[epoch_id]
    bto_us = float(case["bto_us"])
    ges_range = float(np.linalg.norm(satellite_ecef_km - GES_ECEF_KM))
    target = C_KM_S * (bto_us - BTO_BIAS_US) * 1e-6 / 2.0 - ges_range

    south, north = MAP_EXTENT[2], MAP_EXTENT[3]
    # Match the retained BtoArc discretization used by both the public-core and
    # known-flight analyses: 14,001 latitude samples over the declared domain.
    latitude = np.linspace(south, north, ARC_LATITUDE_GRID_POINTS)
    phi = np.radians(latitude)
    prime_vertical = WGS84_A_KM / np.sqrt(1.0 - WGS84_E2 * np.sin(phi) ** 2)
    radius_xy = (prime_vertical + ARC_ALTITUDE_KM) * np.cos(phi)
    z = (prime_vertical * (1.0 - WGS84_E2) + ARC_ALTITUDE_KM) * np.sin(phi)
    satellite_xy = math.hypot(satellite_ecef_km[0], satellite_ecef_km[1])
    satellite_lon = math.atan2(satellite_ecef_km[1], satellite_ecef_km[0])
    aircraft_radius_squared = radius_xy ** 2 + z ** 2
    right = (
        float(np.dot(satellite_ecef_km, satellite_ecef_km))
        + aircraft_radius_squared - target ** 2
    ) / 2.0 - satellite_ecef_km[2] * z
    cosine_delta = right / (radius_xy * satellite_xy)
    valid = np.isfinite(cosine_delta) & (cosine_delta >= -1.0) & (cosine_delta <= 1.0)
    latitude = latitude[valid]
    longitude = np.degrees(satellite_lon + np.arccos(cosine_delta[valid]))
    longitude = (longitude + 180.0) % 360.0 - 180.0
    in_domain = (longitude >= MAP_EXTENT[0]) & (longitude <= MAP_EXTENT[1])
    latitude = latitude[in_domain]
    longitude = longitude[in_domain]
    if len(latitude) < 2:
        raise MaterializationError(f"{epoch_id}: eastern BTO branch misses map domain")
    vectors = unit_vectors(latitude, longitude)
    return Arc(
        epoch_id=epoch_id,
        handshake=pd.Timestamp(str(case["handshake"])),
        bto_us=bto_us,
        satellite_ecef_km=satellite_ecef_km,
        target_range_km=target,
        latitude_deg=latitude,
        longitude_deg_e=longitude,
        unit_vectors=vectors,
        index=cKDTree(vectors),
    )


def physical_slots(handshake: pd.Timestamp) -> tuple[pd.Timestamp, ...]:
    candidates = pd.date_range(
        handshake.floor("2min") - pd.Timedelta(minutes=12),
        handshake.ceil("2min") + pd.Timedelta(minutes=12),
        freq="2min",
    )
    selected = tuple(
        value for value in candidates
        if abs((
            value + pd.Timedelta(seconds=WSPR_MIDPOINT_OFFSET_SECONDS) - handshake
        ).total_seconds()) <= WINDOW_MINUTES * 60.0
    )
    if len(selected) != 10:
        raise MaterializationError(
            f"{handshake}: expected ten physical WSPR slots, got {len(selected)}"
        )
    return selected


def read_wspr(path: Path) -> pd.DataFrame:
    spots = pd.read_csv(path)
    missing = set(REQUIRED_WSPR_COLUMNS).difference(spots.columns)
    if missing:
        raise MaterializationError(f"WSPR input lacks columns {sorted(missing)}")
    spots = spots.loc[spots["band"].isin(INCLUDED_BANDS)].copy()
    spots["time"] = pd.to_datetime(spots["time"], errors="raise")
    spots.sort_values(["time", "id"], kind="mergesort", inplace=True)
    spots.reset_index(drop=True, inplace=True)
    if spots["id"].duplicated().any():
        raise MaterializationError("WSPR input has duplicate report ids")
    return spots


def validate_canonical_observations(path: Path) -> dict[str, object]:
    observations = pd.read_csv(path)
    required = {"epoch_id", "time_utc", "bto_us", "raw_source_rows", "note"}
    missing = required.difference(observations.columns)
    if missing:
        raise MaterializationError(
            f"canonical SATCOM observations lack columns {sorted(missing)}"
        )
    by_epoch = observations.set_index("epoch_id", drop=False)
    selected_rows = []
    for epoch_id, case in EPOCHS.items():
        if epoch_id not in by_epoch.index:
            raise MaterializationError(f"canonical SATCOM input lacks {epoch_id}")
        row = by_epoch.loc[epoch_id]
        if isinstance(row, pd.DataFrame):
            raise MaterializationError(f"canonical SATCOM input duplicates {epoch_id}")
        actual_time = pd.Timestamp(row["time_utc"]).tz_localize(None)
        expected_time = pd.Timestamp(str(case["handshake"]))
        if actual_time != expected_time:
            raise MaterializationError(
                f"{epoch_id}: configured time {expected_time} != canonical {actual_time}"
            )
        if not math.isclose(float(row["bto_us"]), float(case["bto_us"]), abs_tol=1e-12):
            raise MaterializationError(
                f"{epoch_id}: configured BTO {case['bto_us']} != canonical {row['bto_us']}"
            )
        selected_rows.append({
            "epoch_id": epoch_id,
            "time_utc": str(actual_time),
            "bto_us": float(row["bto_us"]),
            "raw_source_rows": str(row["raw_source_rows"]),
            "note": str(row["note"]),
        })
    return {
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "selected_rows": selected_rows,
    }


def validate_canonical_satellite(
    path: Path, states: dict[int, np.ndarray]
) -> dict[str, object]:
    frame = pd.read_csv(path)
    required = {"epoch_id", "time_utc", "x_km", "y_km", "z_km"}
    missing = required.difference(frame.columns)
    if missing:
        raise MaterializationError(
            f"canonical satellite ephemeris lacks columns {sorted(missing)}"
        )
    by_epoch = frame.set_index("epoch_id", drop=False)
    selected_rows = []
    for epoch_id, case in EPOCHS.items():
        if epoch_id not in by_epoch.index:
            raise MaterializationError(f"canonical satellite input lacks {epoch_id}")
        row = by_epoch.loc[epoch_id]
        if isinstance(row, pd.DataFrame):
            raise MaterializationError(f"canonical satellite input duplicates {epoch_id}")
        actual_time = pd.Timestamp(row["time_utc"]).tz_localize(None)
        expected_time = pd.Timestamp(str(case["handshake"]))
        if actual_time != expected_time:
            raise MaterializationError(
                f"{epoch_id}: satellite time {actual_time} != configured {expected_time}"
            )
        workbook_state = states[int(case["ephemeris_second"])]
        canonical_state = row[["x_km", "y_km", "z_km"]].to_numpy(dtype=float)
        if not np.allclose(workbook_state, canonical_state, rtol=0.0, atol=1e-9):
            raise MaterializationError(
                f"{epoch_id}: workbook and canonical satellite ECEF states differ"
            )
        selected_rows.append({
            "epoch_id": epoch_id,
            "time_utc": str(actual_time),
            "x_km": canonical_state[0],
            "y_km": canonical_state[1],
            "z_km": canonical_state[2],
        })
    return {
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "selected_rows": selected_rows,
    }


def _window_z_scores(spots: pd.DataFrame, field: str) -> np.ndarray:
    result = np.full(len(spots), np.nan, dtype=float)
    all_times = spots["time"].to_numpy(dtype="datetime64[ns]").astype(np.int64)
    all_ids = spots["id"].to_numpy(dtype=np.int64)
    all_values = spots[field].to_numpy(dtype=float)
    window_ns = pd.Timedelta(hours=BASELINE_HOURS).value
    for raw_indices in spots.groupby(
        ["tx_sign", "rx_sign", "band"], sort=False, dropna=False
    ).indices.values():
        indices = np.asarray(raw_indices, dtype=np.int64)
        indices = indices[np.lexsort((all_ids[indices], all_times[indices]))]
        times = all_times[indices]
        values = all_values[indices]
        left = np.searchsorted(times, times - window_ns, side="left")
        right = np.searchsorted(times, times + window_ns, side="right")
        count = right - left

        # Center before forming cumulative squares to avoid cancellation for Hz.
        center = float(values.mean())
        delta = values - center
        cumulative = np.concatenate(([0.0], np.cumsum(delta)))
        cumulative_square = np.concatenate(([0.0], np.cumsum(delta * delta)))
        total = cumulative[right] - cumulative[left]
        total_square = cumulative_square[right] - cumulative_square[left]
        mean = center + total / count
        numerator = total_square - total * total / count
        variance = numerator / np.maximum(count - 1, 1)
        # A constant sub-window can leave a tiny positive cancellation residue
        # when the full group mean differs. Treat round-off-scale numerators as
        # zero rather than manufacturing an arbitrarily large z-score.
        cancellation_scale = np.maximum(
            total_square, total * total / np.maximum(count, 1)
        )
        variance_roundoff = (
            64.0 * np.finfo(float).eps * cancellation_scale
        ) / np.maximum(count - 1, 1)
        valid = (
            (count >= MINIMUM_BASELINE_ROWS)
            & (variance > variance_roundoff)
        )
        local = np.full(len(indices), np.nan, dtype=float)
        local[valid] = (
            np.abs(values[valid] - mean[valid]) / np.sqrt(variance[valid])
        )
        result[indices] = local
    return result


def classify_anomalies(spots: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    tx_valid = spots["tx_loc"].fillna("").astype(str).map(
        lambda value: LOCATOR_PATTERN.fullmatch(value) is not None
    )
    rx_valid = spots["rx_loc"].fillna("").astype(str).map(
        lambda value: LOCATOR_PATTERN.fullmatch(value) is not None
    )
    z_snr = _window_z_scores(spots, "snr")
    z_frequency = _window_z_scores(spots, "frequency")
    snr_anomaly = z_snr >= ANOMALY_THRESHOLD
    frequency_anomaly = z_frequency >= ANOMALY_THRESHOLD
    selected = (tx_valid & rx_valid).to_numpy() & (snr_anomaly | frequency_anomaly)
    anomalies = spots.loc[selected].copy()
    anomalies["z_snr"] = z_snr[selected]
    anomalies["z_frequency"] = z_frequency[selected]
    anomalies["type"] = np.select(
        (snr_anomaly[selected] & frequency_anomaly[selected], snr_anomaly[selected]),
        ("dual", "snr_only"),
        default="frequency_only",
    )

    decoded = {}
    for locator in pd.unique(pd.concat((
        anomalies["tx_loc"].str.upper(), anomalies["rx_loc"].str.upper()
    ), ignore_index=True)):
        decoded[str(locator)] = maidenhead_six_character_center(str(locator))
    anomalies["tx_lat_center"] = anomalies["tx_loc"].str.upper().map(
        lambda value: decoded[value][0]
    )
    anomalies["tx_lon_center"] = anomalies["tx_loc"].str.upper().map(
        lambda value: decoded[value][1]
    )
    anomalies["rx_lat_center"] = anomalies["rx_loc"].str.upper().map(
        lambda value: decoded[value][0]
    )
    anomalies["rx_lon_center"] = anomalies["rx_loc"].str.upper().map(
        lambda value: decoded[value][1]
    )
    anomalies.sort_values(["time", "id"], kind="mergesort", inplace=True)
    audit = {
        "input_rows_in_included_bands": len(spots),
        "six_character_geometry_rows": int((tx_valid & rx_valid).sum()),
        "anomalous_six_character_rows": len(anomalies),
        "anomaly_type_counts": {
            str(key): int(value) for key, value in anomalies["type"].value_counts().items()
        },
    }
    return anomalies, audit


def baseline_coverage_audit(spots: pd.DataFrame) -> dict[str, object]:
    radio_times = [
        slot + pd.Timedelta(minutes=offset)
        for case in EPOCHS.values()
        for slot in physical_slots(pd.Timestamp(str(case["handshake"])))
        for offset in CONTROL_OFFSETS_MINUTES.values()
    ]
    selected_minimum = min(radio_times)
    selected_maximum = max(radio_times)
    required_minimum = selected_minimum - pd.Timedelta(hours=BASELINE_HOURS)
    required_maximum = selected_maximum + pd.Timedelta(hours=BASELINE_HOURS)
    input_minimum = pd.Timestamp(spots["time"].min())
    input_maximum = pd.Timestamp(spots["time"].max())
    if input_minimum > required_minimum or input_maximum < required_maximum:
        raise MaterializationError(
            "WSPR input does not cover the full plus/minus-three-hour baseline "
            f"for selected radio slots: input={input_minimum}..{input_maximum}, "
            f"required={required_minimum}..{required_maximum}"
        )
    return {
        "selected_radio_time_min": str(selected_minimum),
        "selected_radio_time_max": str(selected_maximum),
        "required_baseline_time_min": str(required_minimum),
        "required_baseline_time_max": str(required_maximum),
        "input_time_min": str(input_minimum),
        "input_time_max": str(input_maximum),
        "leading_margin_minutes": (
            required_minimum - input_minimum
        ).total_seconds() / 60.0,
        "trailing_margin_minutes": (
            input_maximum - required_maximum
        ).total_seconds() / 60.0,
        "full_baseline_coverage": True,
    }


def all_intersections(anomalies: pd.DataFrame) -> pd.DataFrame:
    """Return both antipodal intersections in the broad map domain."""
    if len(anomalies) < 2:
        return pd.DataFrame()
    ordered = anomalies.sort_values("id", kind="mergesort").reset_index(drop=True)
    transmitter = unit_vectors(ordered["tx_lat_center"], ordered["tx_lon_center"])
    receiver = unit_vectors(ordered["rx_lat_center"], ordered["rx_lon_center"])
    normals = np.cross(transmitter, receiver)
    normal_lengths = np.linalg.norm(normals, axis=1)
    valid_link = normal_lengths > 1e-12
    if not valid_link.all():
        ordered = ordered.loc[valid_link].reset_index(drop=True)
        normals = normals[valid_link]
        normal_lengths = normal_lengths[valid_link]
    normals /= normal_lengths[:, None]
    first, second = np.triu_indices(len(ordered), 1)
    dot = np.einsum("ij,ij->i", normals[first], normals[second])
    angles = np.degrees(np.arccos(np.clip(np.abs(dot), 0.0, 1.0)))
    angle_valid = angles >= MINIMUM_CROSSING_ANGLE_DEG
    first, second, angles = first[angle_valid], second[angle_valid], angles[angle_valid]
    points = np.cross(normals[first], normals[second])
    lengths = np.linalg.norm(points, axis=1)
    valid = lengths > 1e-12
    first, second, angles, points, lengths = (
        first[valid], second[valid], angles[valid], points[valid], lengths[valid]
    )
    points /= lengths[:, None]
    points = np.concatenate((points, -points), axis=0)
    first = np.concatenate((first, first))
    second = np.concatenate((second, second))
    angles = np.concatenate((angles, angles))
    latitude = np.degrees(np.arcsin(np.clip(points[:, 2], -1.0, 1.0)))
    longitude = np.degrees(np.arctan2(points[:, 1], points[:, 0]))
    west, east, south, north = MAP_EXTENT
    in_domain = (
        (longitude >= west) & (longitude <= east)
        & (latitude >= south) & (latitude <= north)
    )
    first, second, angles = first[in_domain], second[in_domain], angles[in_domain]
    latitude, longitude = latitude[in_domain], longitude[in_domain]
    if not len(first):
        return pd.DataFrame()
    return pd.DataFrame({
        "spot_id_1": ordered["id"].to_numpy(dtype=np.int64)[first],
        "spot_id_2": ordered["id"].to_numpy(dtype=np.int64)[second],
        "tx_1": ordered["tx_sign"].to_numpy()[first],
        "rx_1": ordered["rx_sign"].to_numpy()[first],
        "band_1": ordered["band"].to_numpy()[first],
        "type_1": ordered["type"].to_numpy()[first],
        "tx_2": ordered["tx_sign"].to_numpy()[second],
        "rx_2": ordered["rx_sign"].to_numpy()[second],
        "band_2": ordered["band"].to_numpy()[second],
        "type_2": ordered["type"].to_numpy()[second],
        "crossing_angle_deg": angles,
        "lat_deg": latitude,
        "lon_deg_E": longitude,
    })


def filter_to_arc(intersections: pd.DataFrame, arc: Arc) -> pd.DataFrame:
    if intersections.empty:
        return intersections.copy()
    vectors = unit_vectors(intersections["lat_deg"], intersections["lon_deg_E"])
    chord, nearest = arc.index.query(vectors, workers=1)
    distance = EARTH_MEAN_RADIUS_KM * 2.0 * np.arcsin(np.clip(chord / 2.0, 0.0, 1.0))
    retained = distance <= BAND_NM * 1.852
    result = intersections.loc[retained].copy()
    nearest = nearest[retained]
    distance = distance[retained]
    signed = np.where(
        result["lon_deg_E"].to_numpy() >= arc.longitude_deg_e[nearest], distance, -distance
    )
    result["distance_to_arc_km"] = distance
    result["signed_arc_distance_nm"] = signed / 1.852
    result["arc_side"] = np.where(signed >= 0.0, "outward_east", "inward_west")
    return result


def public_core_cluster_audit(pairs: pd.DataFrame) -> dict[str, object]:
    """Reapply the pre-PHaRLAP clustering/residual contract as a checksum."""
    candidates = tuple(
        control.CandidatePair(
            candidate_id=f"post-restart:{row.epoch}:{index:08d}",
            epoch_id=str(row.epoch),
            condition=str(row.condition),
            physical_slot=str(row.physical_slot),
            radio_time=str(row.radio_time),
            spot_id_1=int(row.spot_id_1),
            spot_id_2=int(row.spot_id_2),
            transmitter_1=str(row.tx_1), receiver_1=str(row.rx_1),
            transmitter_2=str(row.tx_2), receiver_2=str(row.rx_2),
            latitude_deg=float(row.lat_deg),
            longitude_deg_e=float(row.lon_deg_E),
            crossing_angle_deg=float(row.crossing_angle_deg),
            distance_to_arc_km=float(row.distance_to_arc_km),
            arc_side=str(row.arc_side),
        )
        for index, row in enumerate(pairs.itertuples(index=False), 1)
    )
    screened = tuple(
        control.ScreenedPair(pair, control.SURFACE_SCREEN, None, {})
        for pair in candidates
    )
    clusters = control.cluster_screened_pairs(screened, 25.0, 3)
    residuals = control.symmetric_residuals(clusters, 25.0)
    counts: dict[str, object] = {}
    for epoch_id in EPOCHS:
        counts[epoch_id] = {
            "raw_clusters": {
                condition: sum(
                    cluster.epoch_id == epoch_id and cluster.condition == condition
                    for cluster in clusters
                )
                for condition in CONTROL_OFFSETS_MINUTES
            },
            "residual_clusters": {
                condition: sum(
                    item.is_residual
                    and item.cluster.epoch_id == epoch_id
                    and item.cluster.condition == condition
                    for item in residuals
                )
                for condition in CONTROL_OFFSETS_MINUTES
            },
        }
    return {
        "purpose": (
            "Deterministic pre-PHaRLAP checksum only; these are not the extended-"
            "model results. Every candidate is provisionally passed through the "
            "surface screen before 25-km clustering and symmetric residualization."
        ),
        "total_clusters": len(clusters),
        "counts": counts,
    }


def materialize(
    wspr_path: Path,
    ephemeris_path: Path,
    canonical_observations_path: Path,
    canonical_satellite_path: Path,
    wspr_source_url: str,
    wspr_query: str,
    wspr_retrieved_utc: str,
    ephemeris_source_url: str,
) -> tuple[bytes, bytes, bytes, bytes, dict[str, object]]:
    canonical_observations = validate_canonical_observations(
        canonical_observations_path
    )
    states = read_ephemeris(ephemeris_path)
    canonical_satellite = validate_canonical_satellite(
        canonical_satellite_path, states
    )
    spots = read_wspr(wspr_path)
    baseline_coverage = baseline_coverage_audit(spots)
    anomalies, anomaly_audit = classify_anomalies(spots)
    arcs = {
        epoch_id: construct_arc(epoch_id, states[int(case["ephemeris_second"])])
        for epoch_id, case in EPOCHS.items()
    }
    grouped = {
        pd.Timestamp(time): frame.copy()
        for time, frame in anomalies.groupby("time", sort=False)
    }
    intersection_cache: dict[pd.Timestamp, pd.DataFrame] = {}
    pair_frames: list[pd.DataFrame] = []
    audit_rows: list[dict[str, object]] = []
    pair_counts: dict[str, dict[str, int]] = {}

    for epoch_id, case in EPOCHS.items():
        arc = arcs[epoch_id]
        slots = physical_slots(arc.handshake)
        pair_counts[epoch_id] = {}
        condition_counts = {name: 0 for name in CONTROL_OFFSETS_MINUTES}
        for condition, offset_minutes in CONTROL_OFFSETS_MINUTES.items():
            for physical_slot in slots:
                radio_time = physical_slot + pd.Timedelta(minutes=offset_minutes)
                if radio_time not in intersection_cache:
                    intersection_cache[radio_time] = all_intersections(
                        grouped.get(radio_time, anomalies.iloc[0:0])
                    )
                frame = filter_to_arc(intersection_cache[radio_time], arc)
                if frame.empty:
                    continue
                frame.insert(0, "radio_time", str(radio_time))
                frame.insert(0, "physical_slot", str(physical_slot))
                frame.insert(0, "condition", condition)
                frame.insert(0, "epoch", epoch_id)
                pair_frames.append(frame)
                condition_counts[condition] += len(frame)
            pair_counts[epoch_id][condition] = condition_counts[condition]

        audit: dict[str, object] = {
            "epoch": epoch_id,
            "sequence": int(case["sequence"]),
            "flight": "MH370 post-radar",
            "title": str(case["title"]),
            "representative_time_utc": str(arc.handshake),
            "bto_us_used": arc.bto_us,
            "raw_bto_us": float(case.get("raw_bto_us", arc.bto_us)),
            "r600_correction_us": float(case.get("r600_correction_us", 0.0)),
            "bto_provenance": str(case["bto_provenance"]),
            "ephemeris_second": int(case["ephemeris_second"]),
            "satellite_x_km": arc.satellite_ecef_km[0],
            "satellite_y_km": arc.satellite_ecef_km[1],
            "satellite_z_km": arc.satellite_ecef_km[2],
            "target_satellite_aircraft_range_km": arc.target_range_km,
            "arc_assumed_aircraft_altitude_km": ARC_ALTITUDE_KM,
            "arc_latitude_min_deg": arc.latitude_deg.min(),
            "arc_latitude_max_deg": arc.latitude_deg.max(),
            "arc_longitude_min_deg_E": arc.longitude_deg_e.min(),
            "arc_longitude_max_deg_E": arc.longitude_deg_e.max(),
            "map_west_deg_E": MAP_EXTENT[0],
            "map_east_deg_E": MAP_EXTENT[1],
            "map_south_deg": MAP_EXTENT[2],
            "map_north_deg": MAP_EXTENT[3],
            "arc_band_each_side_nm": BAND_NM,
            "physical_wspr_slots": ";".join(str(value) for value in slots),
            "physical_wspr_slot_count": len(slots),
            "truth_used_in_candidate_selection": False,
            "post_radar_truth_available": False,
        }
        for condition, offset_minutes in CONTROL_OFFSETS_MINUTES.items():
            audit[f"{condition}_radio_window_start"] = str(
                slots[0] + pd.Timedelta(minutes=offset_minutes)
            )
            audit[f"{condition}_radio_window_end"] = str(
                slots[-1] + pd.Timedelta(minutes=offset_minutes)
            )
            audit[f"{condition}_candidate_pairs"] = condition_counts[condition]
        audit_rows.append(audit)

    if not pair_frames:
        raise MaterializationError("no post-restart candidate pairs were found")
    pairs = pd.concat(pair_frames, ignore_index=True)
    pairs["_epoch_order"] = pairs["epoch"].map(
        {name: index for index, name in enumerate(EPOCHS)}
    )
    pairs["_condition_order"] = pairs["condition"].map(CONDITION_ORDER)
    pairs.sort_values(
        ["_epoch_order", "_condition_order", "physical_slot", "radio_time",
         "spot_id_1", "spot_id_2", "lat_deg", "lon_deg_E"],
        kind="mergesort", inplace=True,
    )
    pairs.drop(columns=["_epoch_order", "_condition_order"], inplace=True)
    pairs.reset_index(drop=True, inplace=True)

    used_ids = set(map(int, pairs["spot_id_1"])) | set(map(int, pairs["spot_id_2"]))
    compact_spots = spots.loc[spots["id"].isin(used_ids), SPOT_FIELDS].copy()
    compact_spots.sort_values("id", kind="mergesort", inplace=True)
    compact_spots.reset_index(drop=True, inplace=True)
    if set(map(int, compact_spots["id"])) != used_ids:
        raise MaterializationError("compact spot table omitted referenced reports")

    arc_rows = []
    for epoch_id, arc in arcs.items():
        case = EPOCHS[epoch_id]
        arc_rows.extend({
            "epoch_id": epoch_id,
            "sequence": int(case["sequence"]),
            "flight": "MH370 post-radar",
            "arc_time_utc": str(arc.handshake),
            "arc_point_index": index,
            "latitude_deg": latitude,
            "longitude_deg_e": longitude,
        } for index, (latitude, longitude) in enumerate(zip(
            arc.latitude_deg, arc.longitude_deg_e
        )))

    public_core_audit = public_core_cluster_audit(pairs)
    pair_payload = deterministic_gzip(dataframe_csv_bytes(pairs))
    spot_payload = deterministic_gzip(dataframe_csv_bytes(compact_spots))
    audit_payload = dataframe_csv_bytes(pd.DataFrame(audit_rows))
    arc_payload = dataframe_csv_bytes(pd.DataFrame(arc_rows))
    manifest: dict[str, object] = {
        "schema": SCHEMA,
        "method": (
            "Trajectory-blind, same-slot public-core anomaly intersections on the "
            "eastern BTO branch, followed by PHaRLAP in a separate frozen batch."
        ),
        "epochs": list(EPOCHS),
        "prespecified_assessment": {
            "map_extent_west_east_south_north_deg": list(MAP_EXTENT),
            "arc_band_each_side_nm": BAND_NM,
            "wspr_midpoint_window_each_side_minutes": WINDOW_MINUTES,
            "control_offsets_minutes": CONTROL_OFFSETS_MINUTES,
            "baseline_each_side_hours": BASELINE_HOURS,
            "minimum_baseline_rows": MINIMUM_BASELINE_ROWS,
            "anomaly_threshold_standard_deviations": ANOMALY_THRESHOLD,
            "minimum_link_circle_crossing_angle_deg": MINIMUM_CROSSING_ANGLE_DEG,
            "included_wspr_bands": list(INCLUDED_BANDS),
            "bto_arc_aircraft_altitude_km": ARC_ALTITUDE_KM,
            "bto_arc_latitude_grid_points": ARC_LATITUDE_GRID_POINTS,
            "truth_used_in_candidate_or_propagation_selection": False,
            "epoch_selection": (
                "One primary BTO arc at each of the seven conventional post-restart "
                "epochs. m1828a/m1828b are additional measurements of the first "
                "restart event, m0019b is the second seventh-arc measurement, and "
                "m1839/m2315 contain BFO but no BTO; none is treated as an additional "
                "numbered arc."
            ),
            "legacy_reproduction_difference": (
                "The earlier exploratory WSPR plot used 18:25:27 and BTO 17120 us "
                "for its first curve. This release assessment instead uses canonical "
                "estimator observation m1825 at 18:25:34 with BTO 12600 us. Other "
                "epoch differences are rounding at approximately one second."
            ),
        },
        "pair_counts": pair_counts,
        "total_pairs": len(pairs),
        "compact_spot_records": len(compact_spots),
        "anomaly_audit": anomaly_audit,
        "baseline_coverage_audit": baseline_coverage,
        "public_core_cluster_audit": public_core_audit,
        "inputs": {
            "wspr_extract": {
                "path_at_materialization": str(wspr_path.resolve()),
                "sha256": sha256_file(wspr_path),
                "source_url": wspr_source_url,
                "query": wspr_query,
                "retrieved_utc": wspr_retrieved_utc,
                "included_band_rows": len(spots),
                "time_min": str(spots["time"].min()),
                "time_max": str(spots["time"].max()),
                "note": "WSPR Live is mutable; the compact frozen output is the PHaRLAP input.",
            },
            "satellite_ephemeris_workbook": {
                "path_at_materialization": str(ephemeris_path.resolve()),
                "sha256": sha256_file(ephemeris_path),
                "source_url": ephemeris_source_url,
            },
            "canonical_satcom_observations": canonical_observations,
            "canonical_satellite_ephemeris": canonical_satellite,
        },
        "outputs": {},
        "runtime": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scipy": __import__("scipy").__version__,
            "openpyxl": __import__("openpyxl").__version__,
        },
        "source_sha256": {
            "materializer": sha256_file(Path(__file__)),
            "control": sha256_file(Path(control.__file__)),
        },
    }
    for name, payload, rows in (
        ("post_restart_pair_locations.csv.gz", pair_payload, len(pairs)),
        ("post_restart_spot_records.csv.gz", spot_payload, len(compact_spots)),
        ("post_restart_arc_audit.csv", audit_payload, len(audit_rows)),
        ("post_restart_arcs.csv", arc_payload, len(arc_rows)),
    ):
        manifest["outputs"][name] = {
            "sha256": hashlib.sha256(payload).hexdigest(), "rows": rows
        }
    return pair_payload, spot_payload, audit_payload, arc_payload, manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--wspr-input", type=Path, required=True)
    parser.add_argument("--ephemeris-workbook", type=Path, required=True)
    parser.add_argument("--canonical-observations", type=Path, required=True)
    parser.add_argument("--canonical-satellite", type=Path, required=True)
    parser.add_argument("--wspr-source-url", default="https://db1.wspr.live/")
    parser.add_argument("--wspr-query", required=True)
    parser.add_argument("--wspr-retrieved-utc", required=True)
    parser.add_argument(
        "--ephemeris-source-url",
        default=(
            "https://raw.githubusercontent.com/gmkf7vfyfb-web/"
            "flight-mh370-revisited/e0115e817975d073bdf2b09a428fbce62aeda35c/"
            "downloads/MH370/Inmarsat3F1_23839%20Fixed%20Position%20Velocity.xlsx"
        ),
    )
    parser.add_argument(
        "--output-directory", type=Path,
        default=Path(__file__).resolve().parents[1] / "data" / "post-restart",
    )
    return parser


def main() -> int:
    arguments = build_parser().parse_args()
    pair, spots, audit, arcs, manifest = materialize(
        arguments.wspr_input, arguments.ephemeris_workbook,
        arguments.canonical_observations,
        arguments.canonical_satellite,
        arguments.wspr_source_url, arguments.wspr_query,
        arguments.wspr_retrieved_utc, arguments.ephemeris_source_url,
    )
    output = arguments.output_directory
    paths = {
        "pairs": output / "post_restart_pair_locations.csv.gz",
        "spots": output / "post_restart_spot_records.csv.gz",
        "audit": output / "post_restart_arc_audit.csv",
        "arcs": output / "post_restart_arcs.csv",
        "manifest": output / "post_restart_candidate_manifest.json",
    }
    for path, payload in (
        (paths["pairs"], pair), (paths["spots"], spots),
        (paths["audit"], audit), (paths["arcs"], arcs),
        (paths["manifest"], canonical_json(manifest)),
    ):
        atomic_write(path, payload)
    print(json.dumps({
        "paths": {key: str(value) for key, value in paths.items()},
        "pair_count": manifest["total_pairs"],
        "spot_count": manifest["compact_spot_records"],
        "pair_counts": manifest["pair_counts"],
    }, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
