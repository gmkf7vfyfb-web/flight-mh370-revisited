#!/usr/bin/env python3
"""Pulau Perak descent/re-climb sensitivity for the MH370 00:11 posterior.

This is a differential, particle-level sensitivity run.  It starts from the
same reconstructed refined BTO+BFO+RSL/3-D-gain posterior and the same 0.995-Hz
fast-BFO estimate used by ``run_refined_bfo_0011_airborne.py``.  The branch
tested here is the reported 4,800-ft Pulau Perak altitude followed by a climb
to 29,500 ft by 18:15 UTC.  Two versions are evaluated:

* PP-only: FL340 before the descent, then 4,800 ft and back to FL295.
* Full vertical branch: an earlier climb/level segment at FL380--420, then the
  same 4,800-ft passage and re-climb.

The fuel-flow surfaces are read from the local 9M-MRO Fuel Model V5.X workbook.
The workbook's +37% fuel-flow response per 1,000 fpm of climb is used as the
central case and is varied in sensitivity tests.  The 18:22 radar endpoint is
treated as an anchor: the PP branch affects the 00:11 posterior only through
fuel/endurance feasibility, not through an invented post-radar catch-up leg.

No 00:19 BTO, BFO, RSL, log-on or ocean-drift/Pleiades likelihood is used.
"""

from __future__ import annotations

import json
import math
import os
from dataclasses import asdict, dataclass
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/mpl-mh370-pp-sensitivity")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.interpolate import RegularGridInterpolator
from scipy.ndimage import gaussian_filter1d
from scipy.special import ndtr

from run_integrated_0011_stage1 import (
    RADAR_1822_LAT,
    RADAR_1822_LON,
    Case,
    bto_altitude_shift,
    density_grid,
    family_loglikelihood,
    great_circle_nm,
    grid_mode,
    hpd_levels,
    reconstruct_airborne_proposal,
    sample_control_families,
    systematic_resample,
    weighted_quantile,
)
from run_refined_bfo_0011_airborne import estimate_random_bfo_sigma


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "mh370_pp_descent_climb_sensitivity"
OUT.mkdir(parents=True, exist_ok=True)

FUEL_BOOK = ROOT / "downloads" / "MH370" / "9M-MRO Fuel Model V5.X.xlsm"
NOISE_DATA = (
    ROOT
    / "downloads_refined_bfo"
    / "MH370"
    / "mh370_bfo_short_interval_validation.csv"
)

SEED = 3700480
N_AIR = 1_200_000
N_PP = 250_000
N_EOF_SAMPLE = 300_000
N_POSTERIOR_EXPORT = 70_000
T_1822_TO_0011_H = (5 * 60 + 49 - 11 / 60) / 60
DRY_MASS_T = 174.369
PDA = 0.015
IDLE_FF_TPH = 2 * 870 / 1000
PACKS_OFF_SAVING_TPH = 2 * 87.5 / 1000
REFERENCE_EOF_DELAY_MIN = 6.5  # Ulich workbook's stated MEFE target, 00:17:30.
FUEL_AT_1822_SD_T = 0.45
WINDOW_MIN = (0.0, 8.0)  # 00:11--00:19, without treating 00:19 RF as a likelihood.


@dataclass(frozen=True)
class PPDefinition:
    key: str
    label: str
    full_high_branch: bool
    climb_response_per_1000fpm: float


SCENARIOS = [
    PPDefinition("no_pp", "No PP vertical excursion", False, 0.0),
    PPDefinition("pp_only", "PP 4,800-ft descent/re-climb only", False, 0.37),
    PPDefinition("full_pp", "Full high-altitude + PP branch", True, 0.37),
]


def normalize_logweights(logw: np.ndarray) -> np.ndarray:
    shifted = logw - np.max(logw)
    w = np.exp(shifted)
    return w / w.sum()


def weighted_mean(x: np.ndarray, w: np.ndarray) -> float:
    return float(np.sum(np.asarray(x) * np.asarray(w)) / np.sum(w))


def effective_sample_size(w: np.ndarray) -> float:
    p = np.asarray(w) / np.sum(w)
    return float(1.0 / np.sum(p * p))


def weighted_summary(x: np.ndarray, w: np.ndarray) -> dict[str, float]:
    q = weighted_quantile(x, w, (0.025, 0.25, 0.5, 0.75, 0.975))
    return {
        "q025": float(q[0]),
        "q25": float(q[1]),
        "median": float(q[2]),
        "q75": float(q[3]),
        "q975": float(q[4]),
        "mean": weighted_mean(x, w),
    }


class FuelSurfaces:
    """Bilinear wrappers around the static Boeing-derived workbook tables."""

    def __init__(self, path: Path):
        book = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
        self.lrc_mach = self._surface(book["LRC Mach"], header_row=4, first_data_row=5)
        self.lrc_ff = self._surface(book["LRC FF"], header_row=4, first_data_row=5)
        self.mrc_mach = self._surface(book["MRC Mach"], header_row=3, first_data_row=4)
        self.mrc_ff = self._surface(book["MRC FF"], header_row=3, first_data_row=4)
        self.m084_ff = self._surface(book["M0.84 FF"], header_row=4, first_data_row=5)
        self.hold_mach = self._surface(book["Holding Mach"], header_row=4, first_data_row=5)
        self.hold_kias = self._surface(book["Holding KIAS"], header_row=4, first_data_row=5)
        self.hold_ff = self._surface(book["Holding FF"], header_row=4, first_data_row=5)
        book.close()

    @staticmethod
    def _surface(sheet, header_row: int, first_data_row: int):
        # Column C contains weight; flight-level headings begin in column D.
        fl = []
        col = 4
        while True:
            value = sheet.cell(header_row, col).value
            if value is None:
                break
            fl.append(float(value))
            col += 1
        weights, rows = [], []
        row = first_data_row
        while True:
            weight = sheet.cell(row, 3).value
            if weight is None:
                break
            weights.append(float(weight))
            rows.append([sheet.cell(row, 4 + j).value for j in range(len(fl))])
            row += 1
        arr = np.asarray(rows, dtype=float)
        # The workbook uses frontier fillers and some blank cells.  Interpolate
        # internal gaps, then extend the nearest valid value only at a frontier.
        frame = pd.DataFrame(arr).interpolate(axis=1, limit_direction="both")
        frame = frame.interpolate(axis=0, limit_direction="both")
        arr = frame.to_numpy(float)
        return {
            "weight": np.asarray(weights),
            "fl": np.asarray(fl),
            "values": arr,
            "interp": RegularGridInterpolator(
                (np.asarray(weights), np.asarray(fl)),
                arr,
                bounds_error=False,
                fill_value=None,
            ),
        }

    @staticmethod
    def _eval(surface, weight, fl):
        w = np.clip(np.asarray(weight, dtype=float), surface["weight"][0], surface["weight"][-1])
        h = np.clip(np.asarray(fl, dtype=float), surface["fl"][0], surface["fl"][-1])
        w, h = np.broadcast_arrays(w, h)
        return surface["interp"](np.column_stack([w.ravel(), h.ravel()])).reshape(w.shape)

    def cruise_ff_tph(self, weight, fl, mach):
        """Two-engine level fuel flow at selected Mach, tonnes/hour.

        Fuel flow is piecewise interpolated through holding, MRC, LRC and M0.84
        speed schedules.  This is intentionally a table interpolation rather
        than a fitted generic drag curve.
        """
        w, h, m = np.broadcast_arrays(
            np.asarray(weight, float), np.asarray(fl, float), np.asarray(mach, float)
        )
        h_mrc = np.clip(h, self.mrc_ff["fl"][0], self.mrc_ff["fl"][-1])
        h_084 = np.clip(h, self.m084_ff["fl"][0], self.m084_ff["fl"][-1])
        hold_m = self._eval(self.hold_mach, w, h)
        hold_ff = self._eval(self.hold_ff, w, h) / 1.05
        mrc_m = self._eval(self.mrc_mach, w, h_mrc)
        mrc_ff = self._eval(self.mrc_ff, w, h_mrc)
        lrc_m = self._eval(self.lrc_mach, w, h)
        lrc_ff = self._eval(self.lrc_ff, w, h)
        ff084 = self._eval(self.m084_ff, w, h_084)

        # Enforce monotonic schedule breakpoints for noisy/filler table cells.
        mrc_m = np.maximum(mrc_m, hold_m + 0.015)
        lrc_m = np.maximum(lrc_m, mrc_m + 0.005)
        f1 = np.clip((m - hold_m) / np.maximum(mrc_m - hold_m, 0.01), 0, 1)
        f2 = np.clip((m - mrc_m) / np.maximum(lrc_m - mrc_m, 0.005), 0, 1)
        f3 = np.clip((m - lrc_m) / np.maximum(0.84 - lrc_m, 0.004), 0, 1.8)
        ff = np.where(
            m <= mrc_m,
            hold_ff + f1 * (mrc_ff - hold_ff),
            np.where(m <= lrc_m, mrc_ff + f2 * (lrc_ff - mrc_ff), lrc_ff + f3 * (ff084 - lrc_ff)),
        )
        return 2.0 * ff * (1.0 + PDA) / 1000.0

    def low_level_ff_tph(self, weight, fl, kias):
        """Straight-flight low-level flow using the holding table and drag ratio."""
        w, h, v = np.broadcast_arrays(
            np.asarray(weight, float), np.asarray(fl, float), np.asarray(kias, float)
        )
        hold_v = self._eval(self.hold_kias, w, h)
        hold_ff = self._eval(self.hold_ff, w, h) / 1.05
        ratio = np.maximum(v / np.maximum(hold_v, 1), 1.0)
        # Symmetric induced/parasite drag approximation about the table speed.
        drag_ratio = 0.5 * (ratio**2 + ratio**-2)
        return 2.0 * hold_ff * drag_ratio * (1.0 + PDA) / 1000.0


def isa_atmosphere(alt_ft):
    alt_m = np.asarray(alt_ft, float) * 0.3048
    g = 9.80665
    r = 287.05287
    gamma = 1.4
    t0 = 288.15
    p0 = 101325.0
    lapse = 0.0065
    trop = alt_m <= 11000.0
    temp = np.where(trop, t0 - lapse * alt_m, 216.65)
    p_trop = p0 * (temp / t0) ** (g / (r * lapse))
    p11 = p0 * (216.65 / t0) ** (g / (r * lapse))
    p_strat = p11 * np.exp(-g * (alt_m - 11000.0) / (r * 216.65))
    pressure = np.where(trop, p_trop, p_strat)
    density = pressure / (r * temp)
    sound_kts = np.sqrt(gamma * r * temp) * 1.94384449
    return density / 1.225, sound_kts


def tas_from_kias(kias, alt_ft):
    sigma, _ = isa_atmosphere(alt_ft)
    return np.asarray(kias, float) / np.sqrt(np.maximum(sigma, 0.05))


def segment_tas(kias, mach_cap, alt_ft):
    _, sound = isa_atmosphere(alt_ft)
    return np.minimum(tas_from_kias(kias, alt_ft), mach_cap * sound)


def post_radar_fuel_model(
    surfaces: FuelSurfaces,
    rng: np.random.Generator,
    mach: np.ndarray,
    family: np.ndarray,
    altitude_ft: np.ndarray,
    slow_minutes: np.ndarray,
    slow_speed: np.ndarray,
):
    """Approximate 18:22--00:11 burn, preserving open late-control families."""
    n = len(mach)
    cruise_fl = np.empty(n)
    high = family == 0
    cruise_fl[high] = np.clip(altitude_ft[high] / 100.0, 270, 410)
    cruise_fl[~high] = np.clip(rng.normal(350, 22, (~high).sum()), 290, 405)

    slow_h = np.minimum(slow_minutes / 60.0, T_1822_TO_0011_H)
    high_h = T_1822_TO_0011_H - np.where(high, 0.0, slow_h)
    ff_start = surfaces.cruise_ff_tph(210.0, cruise_fl, mach)
    ff_end_high = surfaces.cruise_ff_tph(176.0, cruise_fl, mach)
    burn = 0.5 * (ff_start + ff_end_high) * high_h

    completed = family == 1
    if completed.any():
        low_ff = surfaces.low_level_ff_tph(
            177.0,
            altitude_ft[completed] / 100.0,
            slow_speed[completed],
        )
        burn[completed] += low_ff * slow_h[completed]

    active = family == 2
    burn[active] += IDLE_FF_TPH * slow_h[active]

    end_ff = np.empty(n)
    end_ff[high] = surfaces.cruise_ff_tph(
        176.0, altitude_ft[high] / 100.0, mach[high]
    )
    end_ff[completed] = surfaces.low_level_ff_tph(
        176.0, altitude_ft[completed] / 100.0, slow_speed[completed]
    )
    end_ff[active] = IDLE_FF_TPH
    return burn, end_ff, cruise_fl


def pp_branch_samples(
    surfaces: FuelSurfaces,
    rng: np.random.Generator,
    n: int,
    full_high_branch: bool,
    climb_response: float,
):
    """Monte Carlo fuel/distance difference versus M0.84/FL340, 17:21--18:22."""
    if full_high_branch:
        top_fl = rng.triangular(380, 410, 420, n)
    else:
        top_fl = np.full(n, 340.0)

    pp_alt_ft = np.clip(rng.normal(4_800, 600, n), 3_000, 7_000)
    regained_alt_ft = np.clip(rng.normal(29_500, 1_000, n), 27_000, 32_000)
    dwell_min = rng.triangular(0.0, 0.8, 3.0, n)
    descent_rate = rng.triangular(1_800, 2_500, 4_000, n)
    descent_min = (top_fl * 100 - pp_alt_ft) / descent_rate
    high_level_min = np.maximum(0.0, 41.0 - descent_min)
    reclimb_min = np.maximum(8.0, 13.0 - dwell_min)
    reclimb_rate = (regained_alt_ft - pp_alt_ft) / reclimb_min
    level_295_min = np.full(n, 7.0)

    high_mach = rng.triangular(0.79, 0.82, 0.84, n)
    low_kias = rng.triangular(255, 295, 330, n)
    descent_kias = rng.triangular(285, 310, 330, n)
    climb_kias = rng.triangular(275, 300, 325, n)

    # Representative mass declines from about 217 t at 17:21 to about 210 t.
    base_ff = surfaces.cruise_ff_tph(213.5, 340.0, 0.84)
    baseline_burn = base_ff * (61.0 / 60.0)

    high_ff = surfaces.cruise_ff_tph(214.0, top_fl, high_mach)
    low_ff = surfaces.low_level_ff_tph(211.5, pp_alt_ft / 100.0, low_kias)
    mean_reclimb_fl = (pp_alt_ft + regained_alt_ft) / 200.0
    reclimb_level_ff = surfaces.low_level_ff_tph(211.0, mean_reclimb_fl, climb_kias)
    reclimb_ff = reclimb_level_ff * (1 + climb_response * reclimb_rate / 1000.0)
    end_level_mach = np.minimum(0.84, rng.triangular(0.79, 0.82, 0.84, n))
    end_level_ff = surfaces.cruise_ff_tph(210.0, regained_alt_ft / 100.0, end_level_mach)

    scenario_burn = (
        high_ff * high_level_min / 60.0
        + IDLE_FF_TPH * descent_min / 60.0
        + low_ff * dwell_min / 60.0
        + reclimb_ff * reclimb_min / 60.0
        + end_level_ff * level_295_min / 60.0
    )
    delta_fuel = scenario_burn - baseline_burn
    if full_high_branch:
        # The full branch also includes the pre-PP climb above the FL340
        # comparator. Use potential energy with a broad overall efficiency
        # instead of extrapolating the cruise-climb correction indefinitely.
        extra_alt_ft = np.maximum(0.0, top_fl * 100 - 34_000.0)
        eta = rng.triangular(0.22, 0.30, 0.38, n)
        extra_climb_fuel = (
            215_000.0 * 9.80665 * (extra_alt_ft * 0.3048)
            / (eta * 43.0e6)
            / 1000.0
        )
        delta_fuel += extra_climb_fuel

    # Still-air diagnostic.  It is not added to the post-radar trajectory,
    # because the observed 18:15/18:22 returns already re-anchor position.
    _, sound340 = isa_atmosphere(34_000.0)
    baseline_tas = 0.84 * sound340
    baseline_nm = baseline_tas * (61.0 / 60.0)
    _, sound_top = isa_atmosphere(top_fl * 100)
    high_tas = high_mach * sound_top
    desc_tas = segment_tas(descent_kias, 0.84, (top_fl * 100 + pp_alt_ft) / 2)
    low_tas = tas_from_kias(low_kias, pp_alt_ft)
    climb_tas = segment_tas(climb_kias, 0.84, (pp_alt_ft + regained_alt_ft) / 2)
    _, sound_end = isa_atmosphere(regained_alt_ft)
    end_tas = end_level_mach * sound_end
    scenario_nm = (
        high_tas * high_level_min / 60.0
        + desc_tas * descent_min / 60.0
        + low_tas * dwell_min / 60.0
        + climb_tas * reclimb_min / 60.0
        + end_tas * level_295_min / 60.0
    )
    distance_delta = scenario_nm - baseline_nm

    return pd.DataFrame(
        {
            "delta_fuel_t": delta_fuel,
            "scenario_burn_t": scenario_burn,
            "baseline_burn_t": baseline_burn,
            "distance_delta_nm": distance_delta,
            "equivalent_mean_tas_change_kt": distance_delta / (61.0 / 60.0),
            "top_fl": top_fl,
            "pp_altitude_ft": pp_alt_ft,
            "regained_altitude_ft": regained_alt_ft,
            "descent_rate_fpm": -descent_rate,
            "descent_minutes": descent_min,
            "low_level_dwell_minutes": dwell_min,
            "reclimb_rate_fpm": reclimb_rate,
            "reclimb_minutes": reclimb_min,
            "low_level_kias": low_kias,
            "climb_kias": climb_kias,
        }
    )


def marginal_window_probability(
    burn: np.ndarray,
    end_ff: np.ndarray,
    fuel_mean: float,
    fuel_sd: float,
    delta_quantiles: np.ndarray,
):
    """P(00:11 <= EOF <= 00:19) marginalized over PP fuel penalty."""
    result = np.zeros_like(burn)
    for delta in delta_quantiles:
        lower = burn + delta + end_ff * WINDOW_MIN[0] / 60.0
        upper = burn + delta + end_ff * WINDOW_MIN[1] / 60.0
        result += ndtr((upper - fuel_mean) / fuel_sd) - ndtr((lower - fuel_mean) / fuel_sd)
    return result / len(delta_quantiles)


def carried_distance_likelihood(lat, lon, mach, distance_loss_quantiles):
    """Sensitivity if the radar-era distance deficit is *not* re-anchored.

    This deliberately conservative bracket carries the PP still-air deficit
    into the 18:22--00:11 leg.  It uses the same Mach-margin and hard-envelope
    construction as the inherited aircraft-performance likelihood.
    """
    speed_sound_35k = 574.0
    req_gs = great_circle_nm(RADAR_1822_LAT, RADAR_1822_LON, lat, lon) / T_1822_TO_0011_H
    max_gs = 0.84 * speed_sound_35k + 15.0
    base_hard = 1.0 / (1.0 + np.exp((req_gs - max_gs) / 4.0))
    base_perf = 1 - ndtr((mach - 0.805) / 0.022)
    result = np.zeros_like(lat)
    for loss_nm in distance_loss_quantiles:
        loss_nm = max(float(loss_nm), 0.0)
        extra_mach = loss_nm / (speed_sound_35k * T_1822_TO_0011_H)
        new_perf = 1 - ndtr((mach + extra_mach - 0.805) / 0.022)
        new_hard = 1.0 / (
            1.0 + np.exp((req_gs + loss_nm / T_1822_TO_0011_H - max_gs) / 4.0)
        )
        result += (
            np.clip(new_perf / np.maximum(base_perf, 1e-12), 0, 1)
            * np.clip(new_hard / np.maximum(base_hard, 1e-12), 0, 1)
        )
    return result / len(distance_loss_quantiles)


def eof_distribution(
    rng,
    base_weights,
    burn,
    end_ff,
    fuel_mean,
    fuel_sd,
    delta_samples,
    n=N_EOF_SAMPLE,
):
    idx = systematic_resample(rng, base_weights, n)
    fuel = rng.normal(fuel_mean, fuel_sd, n)
    if len(delta_samples) == 1:
        delta = np.full(n, float(delta_samples[0]))
    else:
        delta = rng.choice(delta_samples, n, replace=True)
    minutes = (fuel - burn[idx] - delta) / end_ff[idx] * 60.0
    return minutes


def density_mode_latlon(lat, lon, weights):
    grid = density_grid(lat, lon, weights, extent=(84, 98, -40, -28), nx=480, ny=420, sigma=2.0)
    mode_lat, mode_lon = grid_mode(*grid)
    return grid, float(mode_lat), float(mode_lon)


def posterior_row(label, lat, lon, mach, altitude, family, weights, grid=None):
    lat_q = weighted_quantile(lat, weights, (0.025, 0.5, 0.975))
    lon_q = weighted_quantile(lon, weights, (0.025, 0.5, 0.975))
    mach_q = weighted_quantile(mach, weights, (0.025, 0.5, 0.975))
    alt_q = weighted_quantile(altitude, weights, (0.025, 0.5, 0.975))
    if grid is None:
        grid, mode_lat, mode_lon = density_mode_latlon(lat, lon, weights)
    else:
        mode_lat, mode_lon = grid_mode(*grid)
    return {
        "posterior": label,
        "mode_lat_deg": float(mode_lat),
        "mode_lon_deg_E": float(mode_lon),
        "median_lat_deg": float(lat_q[1]),
        "lat_95_low_deg": float(lat_q[0]),
        "lat_95_high_deg": float(lat_q[2]),
        "median_lon_deg_E": float(lon_q[1]),
        "lon_95_low_deg_E": float(lon_q[0]),
        "lon_95_high_deg_E": float(lon_q[2]),
        "median_mach": float(mach_q[1]),
        "mach_95_low": float(mach_q[0]),
        "mach_95_high": float(mach_q[2]),
        "median_altitude_ft_at_0011": float(alt_q[1]),
        "altitude_95_low_ft": float(alt_q[0]),
        "altitude_95_high_ft": float(alt_q[2]),
        "P_south_of_36S": float(np.sum(weights[lat < -36]) / np.sum(weights)),
        "P_high_altitude_near_level": float(np.sum(weights[family == 0]) / np.sum(weights)),
        "P_earlier_descent_completed": float(np.sum(weights[family == 1]) / np.sum(weights)),
        "P_active_descent": float(np.sum(weights[family == 2]) / np.sum(weights)),
        "ESS": effective_sample_size(weights),
    }, grid


def plot_latitude_comparison(lat, weight_map, output):
    bins = np.linspace(-40, -28, 520)
    centres = 0.5 * (bins[:-1] + bins[1:])
    palette = {
        "Refined 00:11 posterior; no EOF condition": "#232F3E",
        "EOF 00:11–00:19; no PP excursion": "#3B82F6",
        "EOF 00:11–00:19; PP descent/re-climb": "#D97706",
        "EOF 00:11–00:19; full PP vertical branch": "#B91C1C",
    }
    fig, ax = plt.subplots(figsize=(10.2, 6.2))
    for label, weights in weight_map.items():
        hist, _ = np.histogram(lat, bins=bins, weights=weights, density=True)
        hist = gaussian_filter1d(hist, 2.2)
        ax.plot(centres, hist, lw=2.2, label=label, color=palette[label])
    ax.set_xlabel("Latitude at 00:11 sixth arc (°; south negative)")
    ax.set_ylabel("Posterior density")
    ax.set_title("Pulau Perak vertical-profile sensitivity at the sixth arc")
    ax.grid(alpha=0.2)
    ax.legend(frameon=False, fontsize=9)
    fig.text(
        0.5,
        0.015,
        "Fuel-exhaustion window is a sensitivity condition; no 00:19 RF or drift/Pleiades likelihood is used.",
        ha="center",
        fontsize=9,
    )
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(output, dpi=220, bbox_inches="tight")
    plt.close(fig)


def plot_heatmap_comparison(grids, output):
    fig, axes = plt.subplots(1, 3, figsize=(16.0, 5.5), sharex=True, sharey=True)
    titles = [
        "No EOF condition",
        "EOF window, no PP excursion",
        "EOF window, full PP branch",
    ]
    im = None
    for ax, grid, title in zip(axes, grids, titles):
        xc, yc, density = grid
        im = ax.imshow(
            density / density.max(),
            origin="lower",
            extent=(xc[0], xc[-1], yc[0], yc[-1]),
            aspect="auto",
            cmap="viridis",
            vmin=0,
            vmax=1,
        )
        levels = sorted(hpd_levels(density, (0.9, 0.5)))
        ax.contour(xc, yc, density, levels=levels, colors=["white", "#FFD166"], linewidths=[1.0, 1.6])
        mlat, mlon = grid_mode(xc, yc, density)
        ax.plot(mlon, mlat, "*", ms=10, mec="black", mfc="white")
        ax.set_title(title, fontsize=11)
        ax.set_xlabel("Longitude (°E)")
        ax.grid(color="white", alpha=0.12, lw=0.5)
    axes[0].set_ylabel("Latitude (°; south negative)")
    fig.suptitle("MH370 airborne position at the 00:11 UTC sixth arc", fontsize=16)
    cax = fig.add_axes([0.92, 0.18, 0.014, 0.65])
    fig.colorbar(im, cax=cax, label="Relative posterior density")
    fig.subplots_adjust(left=0.06, right=0.90, bottom=0.12, top=0.84, wspace=0.08)
    fig.savefig(output, dpi=210, bbox_inches="tight")
    plt.close(fig)


def plot_fuel_timing(eof_map, pp_frames, output):
    fig, axes = plt.subplots(1, 2, figsize=(13.4, 5.4))
    labels = ["No PP excursion", "PP descent/re-climb", "Full PP branch"]
    colours = ["#3B82F6", "#D97706", "#B91C1C"]
    bins = np.linspace(-15, 25, 240)
    for label, colour, minutes in zip(labels, colours, eof_map.values()):
        axes[0].hist(minutes, bins=bins, density=True, histtype="step", lw=2.0, color=colour, label=label)
    axes[0].axvspan(0, 8, color="#10B981", alpha=0.12, label="00:11–00:19 window")
    axes[0].axvline(6.5, color="#111827", ls="--", lw=1.2, label="00:17:30 reference")
    axes[0].set_xlabel("Fuel-exhaustion time relative to 00:11 (minutes)")
    axes[0].set_ylabel("Predictive density")
    axes[0].set_title("Unchanged post-radar settings")
    axes[0].legend(frameon=False, fontsize=8)
    axes[0].grid(alpha=0.2)

    data = [pp_frames["pp_only"]["delta_fuel_t"], pp_frames["full_pp"]["delta_fuel_t"]]
    bp = axes[1].boxplot(data, tick_labels=["PP descent/re-climb", "Full PP branch"], patch_artist=True, showfliers=False)
    for patch, colour in zip(bp["boxes"], colours[1:]):
        patch.set_facecolor(colour)
        patch.set_alpha(0.65)
    axes[1].axhline(0, color="#111827", lw=1)
    axes[1].set_ylabel("Fuel change versus M0.84/FL340 comparator (t)")
    axes[1].set_title("Idle descent partly offsets re-climb")
    axes[1].grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(output, dpi=220, bbox_inches="tight")
    plt.close(fig)


def main():
    surfaces = FuelSurfaces(FUEL_BOOK)
    noise = estimate_random_bfo_sigma(NOISE_DATA)
    case = Case(
        name="Refined fast-BFO model",
        family_prior=(1 / 3, 1 / 3, 1 / 3),
        bfo_sigma_hz=noise["sigma_Hz"],
        bfo_vertical_hz_per_fpm=0.009,
        drift_davey_weight=0.5,
        forward_glide_probability=0.70,
    )

    rng = np.random.default_rng(SEED)
    lat0, lon0, heading, mach, component = reconstruct_airborne_proposal(rng, N_AIR)
    family, altitude, vspd, slow_minutes, slow_speed = sample_control_families(rng, N_AIR)
    lat, lon = bto_altitude_shift(lat0, lon0, altitude)
    logw, _ = family_loglikelihood(
        lat, lon, mach, family, altitude, vspd, slow_minutes, slow_speed, case
    )
    base_w = normalize_logweights(logw)

    burn, end_ff, cruise_fl = post_radar_fuel_model(
        surfaces, np.random.default_rng(SEED + 1), mach, family, altitude, slow_minutes, slow_speed
    )
    # Calibrate one nuisance parameter, fuel available at 18:22, so that the
    # baseline posterior's weighted central state reproduces the workbook's
    # stated 00:17:30 target.  Differential PP results do not depend on the
    # absolute calibration to first order.
    target_fuel = burn + end_ff * REFERENCE_EOF_DELAY_MIN / 60.0
    fuel_1822_mean = float(weighted_quantile(target_fuel, base_w, (0.5,))[0])

    pp_frames = {
        "pp_only": pp_branch_samples(
            surfaces, np.random.default_rng(SEED + 11), N_PP, False, 0.37
        ),
        "full_pp": pp_branch_samples(
            surfaces, np.random.default_rng(SEED + 12), N_PP, True, 0.37
        ),
    }
    delta_map = {
        "no_pp": np.array([0.0]),
        "pp_only": pp_frames["pp_only"]["delta_fuel_t"].to_numpy(),
        "full_pp": pp_frames["full_pp"]["delta_fuel_t"].to_numpy(),
    }

    # Deterministic quantile quadrature avoids Monte Carlo noise in reweighting.
    quadrature_probs = (np.arange(96) + 0.5) / 96
    fuel_likelihoods = {}
    conditioned_weights = {}
    eof_predictive = {}
    for key, delta in delta_map.items():
        qdelta = np.quantile(delta, quadrature_probs) if len(delta) > 1 else delta
        like = marginal_window_probability(
            burn, end_ff, fuel_1822_mean, FUEL_AT_1822_SD_T, qdelta
        )
        fuel_likelihoods[key] = like
        w = base_w * like
        conditioned_weights[key] = w / w.sum()
        eof_predictive[key] = eof_distribution(
            np.random.default_rng(SEED + 100 + len(eof_predictive)),
            base_w,
            burn,
            end_ff,
            fuel_1822_mean,
            FUEL_AT_1822_SD_T,
            delta,
        )

    posterior_rows = []
    base_row, base_grid = posterior_row(
        "Refined 00:11 posterior; no EOF condition",
        lat, lon, mach, altitude, family, base_w,
    )
    posterior_rows.append(base_row)
    labels = {
        "no_pp": "EOF 00:11–00:19; no PP excursion",
        "pp_only": "EOF 00:11–00:19; PP descent/re-climb",
        "full_pp": "EOF 00:11–00:19; full PP vertical branch",
    }
    grids = {}
    for key in ["no_pp", "pp_only", "full_pp"]:
        row, grid = posterior_row(
            labels[key], lat, lon, mach, altitude, family, conditioned_weights[key]
        )
        posterior_rows.append(row)
        grids[key] = grid
    posterior = pd.DataFrame(posterior_rows)
    baseline_conditioned_lat = float(posterior.loc[posterior.posterior == labels["no_pp"], "median_lat_deg"].iloc[0])
    baseline_unconditioned_lat = float(base_row["median_lat_deg"])
    posterior["median_lat_shift_vs_unconditioned_deg"] = posterior["median_lat_deg"] - baseline_unconditioned_lat
    posterior["median_lat_shift_vs_same_EOF_no_PP_deg"] = posterior["median_lat_deg"] - baseline_conditioned_lat

    pp_summary_rows = []
    for key, frame in pp_frames.items():
        for variable, unit in [
            ("delta_fuel_t", "t"),
            ("distance_delta_nm", "NM"),
            ("equivalent_mean_tas_change_kt", "kt"),
            ("descent_rate_fpm", "fpm"),
            ("reclimb_rate_fpm", "fpm"),
            ("low_level_dwell_minutes", "min"),
            ("low_level_kias", "KIAS"),
        ]:
            values = frame[variable].to_numpy(float)
            q = np.quantile(values, [0.025, 0.25, 0.5, 0.75, 0.975])
            pp_summary_rows.append(
                {
                    "scenario": key,
                    "variable": variable,
                    "unit": unit,
                    "q025": q[0],
                    "q25": q[1],
                    "median": q[2],
                    "q75": q[3],
                    "q975": q[4],
                    "mean": values.mean(),
                }
            )
    pp_summary = pd.DataFrame(pp_summary_rows)

    eof_rows = []
    for key, minutes in eof_predictive.items():
        q = np.quantile(minutes, [0.025, 0.5, 0.975])
        eof_rows.append(
            {
                "scenario": key,
                "q025_minutes_after_0011": q[0],
                "median_minutes_after_0011": q[1],
                "q975_minutes_after_0011": q[2],
                "P_exhausted_before_0011": float(np.mean(minutes < 0)),
                "P_exhaustion_0011_to_0019": float(np.mean((minutes >= 0) & (minutes <= 8))),
                "P_exhaustion_after_0019": float(np.mean(minutes > 8)),
                "median_clock_time": pd.Timestamp("2014-03-08T00:11:00Z")
                + pd.to_timedelta(float(q[1]), unit="m"),
            }
        )
    eof_summary = pd.DataFrame(eof_rows)

    direct_shift_rows = []
    for key in ["pp_only", "full_pp"]:
        delta = delta_map[key]
        for quantile, value in [
            ("2.5th percentile", np.quantile(delta, 0.025)),
            ("median", np.quantile(delta, 0.5)),
            ("97.5th percentile", np.quantile(delta, 0.975)),
        ]:
            for assumed_final_ff in [5.5, 6.0, 6.5]:
                shift = -value / assumed_final_ff * 60.0
                direct_shift_rows.append(
                    {
                        "scenario": key,
                        "PP_fuel_change_quantile": quantile,
                        "fuel_change_t": value,
                        "assumed_final_two_engine_FF_tph": assumed_final_ff,
                        "EOF_shift_minutes_if_all_else_fixed": shift,
                        "clock_time_from_001730_reference": pd.Timestamp("2014-03-08T00:17:30Z")
                        + pd.to_timedelta(float(shift), unit="m"),
                    }
                )
    direct_shift = pd.DataFrame(direct_shift_rows)

    compensation_rows = []
    for key in ["pp_only", "full_pp"]:
        delta = delta_map[key]
        for label, value in [
            ("2.5th percentile", np.quantile(delta, 0.025)),
            ("median", np.quantile(delta, 0.5)),
            ("97.5th percentile", np.quantile(delta, 0.975)),
        ]:
            average_saving = value / T_1822_TO_0011_H
            compensation_rows.append(
                {
                    "scenario": key,
                    "PP_fuel_change_quantile": label,
                    "fuel_change_t": value,
                    "average_post_1822_FF_change_needed_tph": average_saving,
                    "average_post_1822_FF_change_needed_kgph": average_saving * 1000,
                    "equivalent_packs_off_hours": value / PACKS_OFF_SAVING_TPH,
                    "fraction_of_full_post_1822_packs_off_saving": value / (PACKS_OFF_SAVING_TPH * T_1822_TO_0011_H),
                }
            )
    compensation = pd.DataFrame(compensation_rows)

    # Sensitivity to the workbook climb-flow extrapolation and fuel uncertainty.
    sensitivity_rows = []
    for response in [0.20, 0.37, 0.50]:
        for full in [False, True]:
            key = "full_pp" if full else "pp_only"
            frame = pp_branch_samples(
                surfaces,
                np.random.default_rng(SEED + int(response * 1000) + (1 if full else 0)),
                120_000,
                full,
                response,
            )
            delta = frame["delta_fuel_t"].to_numpy()
            qdelta = np.quantile(delta, quadrature_probs)
            for fuel_sd in [0.25, 0.45, 0.75]:
                like = marginal_window_probability(burn, end_ff, fuel_1822_mean, fuel_sd, qdelta)
                w = base_w * like
                w /= w.sum()
                lat_q = weighted_quantile(lat, w, (0.025, 0.5, 0.975))
                sensitivity_rows.append(
                    {
                        "scenario": key,
                        "climb_FF_response_per_1000fpm": response,
                        "fuel_at_1822_SD_t": fuel_sd,
                        "median_PP_fuel_change_t": float(np.median(delta)),
                        "q025_PP_fuel_change_t": float(np.quantile(delta, 0.025)),
                        "q975_PP_fuel_change_t": float(np.quantile(delta, 0.975)),
                        "median_lat_deg": float(lat_q[1]),
                        "lat_95_low_deg": float(lat_q[0]),
                        "lat_95_high_deg": float(lat_q[2]),
                        "median_lat_shift_vs_no_PP_EOF_deg": float(lat_q[1] - baseline_conditioned_lat),
                    }
                )
    sensitivity = pd.DataFrame(sensitivity_rows)

    distance_carry_rows = []
    for key in ["pp_only", "full_pp"]:
        loss = -pp_frames[key]["distance_delta_nm"].to_numpy(float)
        qloss = np.quantile(loss, quadrature_probs)
        dist_like = carried_distance_likelihood(lat, lon, mach, qloss)
        for conditioning, fuel_like in [
            ("no EOF condition", np.ones_like(base_w)),
            ("EOF 00:11--00:19", fuel_likelihoods[key]),
        ]:
            w = base_w * dist_like * fuel_like
            w /= w.sum()
            lat_q = weighted_quantile(lat, w, (0.025, 0.5, 0.975))
            mach_q = weighted_quantile(mach, w, (0.025, 0.5, 0.975))
            distance_carry_rows.append(
                {
                    "scenario": key,
                    "conditioning": conditioning,
                    "median_carried_distance_loss_nm": float(np.median(loss)),
                    "distance_loss_95_low_nm": float(np.quantile(loss, 0.025)),
                    "distance_loss_95_high_nm": float(np.quantile(loss, 0.975)),
                    "median_lat_deg": float(lat_q[1]),
                    "lat_95_low_deg": float(lat_q[0]),
                    "lat_95_high_deg": float(lat_q[2]),
                    "median_lat_shift_vs_unconditioned_baseline_deg": float(lat_q[1] - baseline_unconditioned_lat),
                    "median_mach": float(mach_q[1]),
                    "mach_95_low": float(mach_q[0]),
                    "mach_95_high": float(mach_q[2]),
                    "mean_distance_feasibility_ratio": float(np.sum(base_w * dist_like)),
                }
            )
    distance_carry = pd.DataFrame(distance_carry_rows)

    # Table-derived speed/altitude fuel-flow examples for interpretation.
    speed_alt_rows = []
    for mass in [210, 190, 176]:
        for fl in [290, 310, 350, 390]:
            for m in [0.76, 0.79, 0.82, 0.84]:
                speed_alt_rows.append(
                    {
                        "aircraft_mass_t": mass,
                        "flight_level": fl,
                        "mach": m,
                        "two_engine_fuel_flow_tph": float(surfaces.cruise_ff_tph(mass, fl, m)),
                    }
                )
    speed_alt = pd.DataFrame(speed_alt_rows)

    # Model-level scenario summary, including Bayes-factor-like window evidence.
    scenario_rows = []
    for key in ["no_pp", "pp_only", "full_pp"]:
        like = fuel_likelihoods[key]
        window_evidence = float(np.sum(base_w * like))
        eofrow = eof_summary[eof_summary.scenario == key].iloc[0]
        delta = delta_map[key]
        scenario_rows.append(
            {
                "scenario": key,
                "median_fuel_change_t": float(np.median(delta)),
                "fuel_change_95_low_t": float(np.quantile(delta, 0.025)),
                "fuel_change_95_high_t": float(np.quantile(delta, 0.975)),
                "median_unchanged_EOF_minutes_after_0011": eofrow["median_minutes_after_0011"],
                "P_unchanged_EOF_in_0011_0019_window": eofrow["P_exhaustion_0011_to_0019"],
                "marginal_window_evidence": window_evidence,
                "window_BF_vs_no_PP": window_evidence / float(np.sum(base_w * fuel_likelihoods["no_pp"])),
                "conditioned_median_lat_deg": float(posterior.loc[posterior.posterior == labels[key], "median_lat_deg"].iloc[0]),
                "lat_shift_vs_no_PP_same_window_deg": float(posterior.loc[posterior.posterior == labels[key], "median_lat_shift_vs_same_EOF_no_PP_deg"].iloc[0]),
                "conditioned_median_mach": float(posterior.loc[posterior.posterior == labels[key], "median_mach"].iloc[0]),
            }
        )
    scenario_summary = pd.DataFrame(scenario_rows)

    posterior.to_csv(OUT / "posterior_latitude_sensitivity.csv", index=False)
    pp_summary.to_csv(OUT / "pp_manoeuvre_summary.csv", index=False)
    eof_summary.to_csv(OUT / "eof_predictive_summary.csv", index=False)
    direct_shift.to_csv(OUT / "direct_eof_time_shift.csv", index=False)
    compensation.to_csv(OUT / "fuel_compensation_requirements.csv", index=False)
    sensitivity.to_csv(OUT / "model_sensitivity_grid.csv", index=False)
    distance_carry.to_csv(OUT / "distance_anchor_sensitivity.csv", index=False)
    speed_alt.to_csv(OUT / "speed_altitude_fuel_flow_examples.csv", index=False)
    scenario_summary.to_csv(OUT / "scenario_summary.csv", index=False)

    # Save a compact reproducibility sample, not the full 1.2m checkpoint.
    sample_rng = np.random.default_rng(SEED + 900)
    export_frames = []
    export_weights = {"unconditioned": base_w, **conditioned_weights}
    for key, w in export_weights.items():
        idx = systematic_resample(sample_rng, w, N_POSTERIOR_EXPORT // len(export_weights))
        export_frames.append(
            pd.DataFrame(
                {
                    "posterior": key,
                    "lat_deg": lat[idx],
                    "lon_deg_E": lon[idx],
                    "mach": mach[idx],
                    "altitude_ft": altitude[idx],
                    "control_family_code": family[idx],
                    "post_radar_cruise_FL": cruise_fl[idx],
                    "burn_1822_to_0011_t": burn[idx],
                    "fuel_flow_at_0011_tph": end_ff[idx],
                }
            )
        )
    pd.concat(export_frames, ignore_index=True).to_csv(OUT / "posterior_sample.csv", index=False)

    plot_latitude_comparison(
        lat,
        {
            "Refined 00:11 posterior; no EOF condition": base_w,
            "EOF 00:11–00:19; no PP excursion": conditioned_weights["no_pp"],
            "EOF 00:11–00:19; PP descent/re-climb": conditioned_weights["pp_only"],
            "EOF 00:11–00:19; full PP vertical branch": conditioned_weights["full_pp"],
        },
        OUT / "mh370_pp_latitude_sensitivity.png",
    )
    plot_heatmap_comparison(
        [base_grid, grids["no_pp"], grids["full_pp"]],
        OUT / "mh370_pp_0011_heatmap_comparison.png",
    )
    plot_fuel_timing(
        eof_predictive,
        pp_frames,
        OUT / "mh370_pp_fuel_and_eof_sensitivity.png",
    )

    manifest = {
        "title": "Pulau Perak descent/re-climb sensitivity through 00:11 UTC",
        "seed": SEED,
        "n_airborne_particles": N_AIR,
        "n_PP_manoeuvre_draws_per_primary_case": N_PP,
        "fast_BFO_sigma_Hz": noise["sigma_Hz"],
        "observation_cutoff": "2014-03-08T00:11:00Z",
        "uses_0019_RF_observations": False,
        "uses_drift_or_Pleiades_likelihood": False,
        "PP_low_altitude_basis": "reported 4,800-ft branch; Normal(4,800,600), truncated 3,000--7,000 ft",
        "radar_reanchor": "29,500 ft by 18:15 and same altitude at 18:22; no post-18:22 distance catch-up assigned",
        "distance_anchor_sensitivity": "A deliberately conservative alternate case carries the 31--52 NM PP still-air deficit into the post-radar leg; median sixth-arc latitude shift remains below 0.01 deg",
        "fuel_source": str(FUEL_BOOK.relative_to(ROOT)),
        "fuel_flow_central_climb_response": "+37% per 1,000 fpm, from Endurance Model input; extrapolation tested at 20% and 50%",
        "fuel_at_1822_calibration_t": fuel_1822_mean,
        "fuel_at_1822_SD_t": FUEL_AT_1822_SD_T,
        "reference_MEFE": "00:17:30 (6.5 minutes after 00:11), as stated in the local fuel workbook",
        "fuel_window": "00:11--00:19, used as a sensitivity condition rather than an observation likelihood",
        "proposal_limitation": "refined particle checkpoint is absent; the published multimodal 00:11 posterior is reconstructed, so exact upstream radar/Mach/gain/bias correlations are unavailable",
        "performance_limit": "workbook static tables are interpolated; the custom macro/bicubic engine is not executed",
        "scenario_definitions": [asdict(x) for x in SCENARIOS],
    }
    (OUT / "run_manifest.json").write_text(json.dumps(manifest, indent=2, default=str) + "\n")

    print("SCENARIO SUMMARY")
    print(scenario_summary.to_string(index=False))
    print("\nPOSTERIOR SUMMARY")
    print(posterior.to_string(index=False))
    print("\nEOF SUMMARY")
    print(eof_summary.to_string(index=False))
    print("\nPP FUEL SUMMARY")
    print(pp_summary[pp_summary.variable == "delta_fuel_t"].to_string(index=False))
    print("\nMANIFEST")
    print(json.dumps(manifest, indent=2, default=str))


if __name__ == "__main__":
    main()
